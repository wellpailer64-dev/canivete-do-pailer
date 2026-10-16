"""Slider Kanivete — 2ª bateria: janela sem margem, gráficos (todos os modelos, CSV, PPTX nativo), paletas de cliente
(Photoshop, salvar, tema, .ase ida e volta, aplicar), animações e transições novas (prévia na mesa e apresentação),
ferramentas do agente (texto, folha, revisar, achar, lote, substituir, estilo) e TODOS os atalhos com teclado de verdade.
App de teste em --agente=9333.
    py -3.13 testes/teste_slider_extra.py [--porta 9333]
Saídas em %TEMP%/kanivete_teste_slider. Sai com 1 se algo reprovar."""
import argparse, os, sys, tempfile, time
from playwright.sync_api import sync_playwright
sys.stdout.reconfigure(encoding="utf-8")
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
ap = argparse.ArgumentParser(); ap.add_argument("--porta", type=int, default=9333)
a = ap.parse_args()
OUT = os.path.join(tempfile.gettempdir(), "kanivete_teste_slider").replace("\\", "/"); os.makedirs(OUT, exist_ok=True)
falhas = []


def ok(cond, nome, extra=""):
    print(("  ok  " if cond else "  FALHOU ") + nome + (f"  ({extra})" if extra else ""), flush=True)
    if not cond:
        falhas.append(nome)


with sync_playwright() as p:
    pg = next(x for c in p.chromium.connect_over_cdp(f"http://127.0.0.1:{a.porta}").contexts for x in c.pages if "index.html" in x.url)
    pg.set_default_timeout(120000)
    erros = []
    pg.on("pageerror", lambda e: erros.append(str(e)))
    E = lambda js, arg=None: pg.evaluate(js, arg)
    K = lambda tecla: (pg.keyboard.press(tecla), time.sleep(0.12))
    E("switchTool('slider-kanivete')"); time.sleep(0.6)
    E("async () => { await SDK.cmd('modelo', {tema: 'claro', nome: 'Bateria'}); SDK_UI.editor(); }"); time.sleep(0.8)

    print("janela")
    r = E("""() => { const m = document.querySelector('.main-content').getBoundingClientRect(), s = document.getElementById('sd').getBoundingClientRect();
        return {top: s.top - m.top, left: s.left - m.left, pad: getComputedStyle(document.querySelector('.content-area')).paddingTop}; }""")
    ok(r["pad"] == "10px" and r["left"] <= 12, "Slider ocupa a área toda (margem como Photo/Vetor)", str(r))

    print("gráficos")
    ids = E("""async () => { const out = {}; await SDK.cmd('slide_novo', {layout: 'vazio'});
        for (const m of SDK.modelosGrafico()) out[m.id] = (await SDK.grafico({modelo: m.id})).id; return out; }""")
    ok(len(ids) == 14, "14 modelos de gráfico inseridos", str(len(ids)))
    vaz = E("(ids) => Object.entries(ids).filter(([k, id]) => (SDU.mapa.get(id).el.querySelectorAll('rect,path,circle,text').length < 3)).map(x => x[0])", ids)
    ok(not vaz, "todo gráfico desenha na mesa", ",".join(vaz))
    E("(id) => SDK.cmd('grafico_dados', {id, csv: 'Mês;Loja A;Loja B\\nJan;10;20\\nFev;30;5\\nMar;1.234,5;7'})", ids["agrupadas"])
    g = E("(id) => { const c = sdCamada(id); return [c.rotulos, c.series.map(s => s.nome), c.series[0].valores]; }", ids["agrupadas"])
    ok(g[0] == ["Jan", "Fev", "Mar"] and g[1] == ["Loja A", "Loja B"] and abs(g[2][2] - 1234.5) < 0.01, "CSV com número brasileiro vira séries", str(g))
    E("(id) => { SD.sel = [id]; sdAvisar('sel'); }", ids["pizza"]); time.sleep(0.3)
    ok(E("!!document.querySelector('#sd-props [data-graf-csv]')"), "painel do gráfico mostra os dados para editar")
    E("""() => { const t = document.querySelector('#sd-props [data-graf-csv]'); t.value = 'Fatia;Valor\\nA;50\\nB;50'; t.dispatchEvent(new Event('change', {bubbles: true})); }"""); time.sleep(0.3)
    ok(E("(id) => sdCamada(id).rotulos.join() === 'A,B'", ids["pizza"]), "editar a caixa de dados muda o gráfico")
    foto = E("async () => { const d = await SDK.foto(SD.atual + 1, 480); return d.length; }")
    ok(foto > 15000, "foto do slide com 14 gráficos", f"{foto} bytes")
    pptx = f"{OUT}/graficos.pptx"
    if os.path.isfile(pptx):
        os.remove(pptx)
    E("(c) => SDK.cmd('exportar', {formato: 'pptx', caminho: c})", pptx)
    from pptx import Presentation
    n = sum(1 for s in Presentation(pptx).slides for sh in s.shapes if sh.has_chart)
    ok(n == 14, "PPTX editável com gráficos NATIVOS do PowerPoint", f"{n} gráficos")
    d = E("(id) => SDK.cmd('duplicar', {ids: [id]})", ids["kpi"])
    ok(len(d["ids"]) == 1, "duplicar gráfico")

    print("grupos")
    t = E("async () => { await SDK.cmd('slide_novo', {layout: 'vazio'}); return SDK.cmd('tabela', {linhas: 2, colunas: 2}); }")
    dup = E("async (ids) => { const r = await SDK.cmd('duplicar', {ids}); return [sdCamada(ids[0]).grupo, sdCamada(r.ids[0]).grupo, new Set(r.ids.map(i => sdCamada(i).grupo)).size]; }", t["ids"])
    ok(dup[0] != dup[1] and dup[2] == 1, "cópia de grupo vira grupo novo (não seleciona a original junto)", str(dup))

    print("paletas")
    ps = E("async () => (await SDK.cmd('paletas')).paletas")
    ok(len(ps) >= 1, "paletas carregadas (do Photoshop instalado ou básicas)", f"{len(ps)}: " + ", ".join(x["nome"] for x in ps[:6]))
    E("() => SDK.cmd('paleta_salvar', {nome: 'Cliente Teste', cores: ['#0a2540', '#f6f9fc', '#635bff', '#00d4ff', '#425466'], substituir: true})")
    ps2 = E("async () => (await SDK.cmd('paletas')).paletas.find(x => x.nome === 'Cliente Teste')")
    ok(ps2 and ps2["cores"][2] == "#635bff", "salvar paleta do cliente com nome")
    time.sleep(0.6)
    ok(E("async () => { const r = await window.pywebview.api.ie_prefs(); return (r.prefs.amostrasLista || []).some(x => x.g === 'Cliente Teste'); }"),
       "paleta gravada nas preferências (a mesma do Photo)")
    E("() => SDK.cmd('paleta_tema', {nome: 'Cliente Teste'})")
    tm = E("SD.doc.tema.cores")
    ok(tm["fundo"] == "#f6f9fc" and tm["tinta"] == "#0a2540" and tm["destaque"] in ("#635bff", "#00d4ff"), "usar paleta como tema (papéis pelo contraste)", str(tm))
    ase = f"{OUT}/cliente.ase"
    r = E("([p]) => window.pywebview.api.ie_amostras_gravar(p, [{g: 'Cliente Teste', cores: [{c: '#635bff', n: 'Roxo'}, {c: '#00d4ff', n: ''}]}])", [ase])
    r2 = E("(p) => window.pywebview.api.ie_amostras(p)", ase)
    ok(r.get("success") and r2.get("success") and r2["lista"][0]["g"] == "Cliente Teste" and r2["lista"][0]["cores"][0]["c"] == "#635bff",
       "exportar .ase e ler de volta (grupo e cor)")
    E("() => { sdAba('cores'); }"); time.sleep(0.4)
    ok(E("document.querySelectorAll('#sd-cores .sd-am').length > 5 && !document.getElementById('sd-cores').hidden"), "aba Cores mostra as amostras")
    tid = E("async () => (await SDK.cmd('texto', {texto: 'Cor por amostra'})).id")
    E("(id) => { SD.sel = [id]; sdAvisar('sel'); }", tid); time.sleep(0.3)
    pg.click("#sd-cores .sd-am[data-am='#635bff']"); time.sleep(0.3)
    ok(E("(id) => sdCamada(id).cor === '#635bff'", tid), "clicar na amostra pinta o texto selecionado")
    E("() => sdAba('design')")

    print("animações e transições")
    E("(id) => SDK.cmd('animacao', {ids: [id], tipo: 'crescer', dur: 0.8})", tid)
    n = E("(id) => { sdPreverAnim([id]); return SDU.mapa.get(id).el.getAnimations().length; }", tid)
    ok(n >= 1, "prévia da animação toca na mesa", str(n))
    E("() => SDK.cmd('slide_transicao', {tipo: 'revelar', dur: 0.6})")
    ok(E("sdPreverTransicao() && document.querySelector('#sd-palco .sd-slide').getAnimations().length >= 1"), "prévia da transição toca na mesa")
    time.sleep(0.4)
    ok(E("(() => { const f = document.querySelector('.sd-it.on .sd-it-fx'); return f && !f.hidden && f.textContent.includes('⇢') && f.textContent.includes('★'); })()"),
       "miniatura mostra selo de transição e animação")
    E("() => sdApresentar(SD.atual)"); time.sleep(0.3)
    cp = E("(id) => { const e = SDS.el.querySelector(`[data-id='${id}']`); return [getComputedStyle(e).clipPath, e.getAnimations().length]; }", tid)
    time.sleep(1.2)
    cp2 = E("(id) => getComputedStyle(SDS.el.querySelector(`[data-id='${id}']`)).clipPath", tid)
    ok(cp[0].startswith("inset(") and cp[1] >= 1 and cp2 == "none", "apresentação: 'crescer' revela animando e solta o recorte no fim", f"{cp} → {cp2}")
    K("Escape"); time.sleep(0.3)
    for tr in ["subir", "preto", "revelar", "empurrar"]:
        E("(t) => SDK.cmd('slide_transicao', {tipo: t, todos: true, dur: 0.3})", tr)
    E("() => sdApresentar(0)"); time.sleep(0.3); K("ArrowRight"); time.sleep(0.5)
    ok(E("SDS.on && SDS.i === 1"), "apresentação avança com transição nova")
    K("Escape"); time.sleep(0.3)

    print("leveza")
    E("() => SDK.cmd('slide_novo', {layout: 'vazio'})")
    E("(c) => SDK.cmd('video', {src: c, w: 640, h: 360, mudo: true})", os.path.join(OUT, "clipe.mp4").replace("\\", "/")) if os.path.isfile(os.path.join(OUT, "clipe.mp4")) else None
    E("() => { window._v = document.querySelector('#sd-palco video'); window._m = document.querySelector('#sd-lista [data-mini]:not([data-mini=\"' + sdSlide().id + '\"]) .sd-slide'); }")
    E("() => SDK.cmd('texto', {texto: 'Comando novo'})")
    ok(E("!window._v || document.querySelector('#sd-palco video') === window._v"), "vídeo da mesa não recarrega quando outro comando muda o slide")
    ok(E("!!window._m && document.contains(window._m)"), "miniaturas dos outros slides são reaproveitadas (não redesenha a lista toda)")

    print("agente")
    txt = E("SDK.texto()")
    ok("#1" in txt and "G " in txt and len(txt) < 12000, "SDK.texto(): roteiro compacto", f"{len(txt)} caracteres")
    r = E("(c) => SDK.folha({caminho: c, largura: 240})", f"{OUT}/folha.jpg")
    ok(r.get("caminho") and os.path.getsize(f"{OUT}/folha.jpg") > 20000, "SDK.folha(): todas as páginas numa imagem", f"{os.path.getsize(OUT + '/folha.jpg') // 1024} KB")
    E("async () => { await SDK.cmd('slide_novo', {layout: 'vazio'}); await SDK.cmd('texto', {texto: 'Texto comprido demais '.repeat(30), w: 400, h: 60, tam: 40, nome: 'Estouro'}); await SDK.cmd('texto', {texto: 'miúdo', tam: 12, x: 100, y: 900, nome: 'Miudo'}); await SDK.cmd('texto', {texto: 'claro', cor: '#f4f4f4', x: 900, y: 900, nome: 'Claro'}); }")
    rv = E("SDK.revisar(SD.atual + 1)")
    tipos = {x["tipo"] for x in rv["problemas"]}
    ok({"estouro", "pequeno", "contraste"} <= tipos, "SDK.revisar() acha estouro, letra pequena e pouco contraste", ",".join(sorted(tipos)))
    ok(len(E("SDK.achar('comprido')")) == 1, "SDK.achar()")
    h0 = E("SD.hist.length")
    E("() => SDK.lote([['texto', {texto: 'L1'}], ['texto', {texto: 'L2'}], ['slide_novo', {}]])")
    ok(E("SD.hist.length") == h0 + 1, "SDK.lote(): vários comandos, um desfazer só")
    E("SDK.desfazer()")
    ok(not E("SDK.achar('L1').length"), "desfazer o lote inteiro")
    E("() => SDK.cmd('substituir', {de: 'Obrigado!', para: 'Valeu!'})")
    ok(len(E("SDK.achar('Valeu!')")) == 1, "substituir texto em todos os slides")

    print("atalhos (teclado de verdade)")
    E("async () => { await SDK.cmd('slide_novo', {layout: 'vazio'}); }")
    pg.mouse.click(*E("() => { const r = document.getElementById('sd-mesa').getBoundingClientRect(); return [r.left + 8, r.top + 8]; }")); time.sleep(0.2)
    n0 = E("SD.doc.slides.length"); K("Control+m"); ok(E("SD.doc.slides.length") == n0 + 1, "Ctrl+M novo slide")
    K("Control+Shift+d"); ok(E("SD.doc.slides.length") == n0 + 2, "Ctrl+Shift+D duplica o slide")
    E("() => { sdHist(); sdSlide().camadas = []; SD.sel = []; sdMudou('tudo'); }")   # slide limpo (trocar layout agora preserva os textos)
    K("t"); tid = E("SD.sel[0]"); ok(E("sdCamada(SD.sel[0])?.tipo === 'texto' && sdSlide().camadas.length === 1"), "T insere texto")
    x0 = E("sdCamada(SD.sel[0]).x"); K("ArrowRight"); K("Shift+ArrowRight"); time.sleep(0.5)
    ok(E("sdCamada(SD.sel[0]).x") == x0 + 11, "setas movem 1 px, Shift+seta 10 px")
    K("Control+d"); ok(E("sdSlide().camadas.length") == 2, "Ctrl+D duplica a camada")
    K("Control+a"); ok(E("SD.sel.length") == 2, "Ctrl+A seleciona tudo")
    K("Control+g"); ok(E("SD.sel.every(i => sdCamada(i).grupo)"), "Ctrl+G agrupa")
    K("Control+Shift+g"); ok(E("SD.sel.every(i => !sdCamada(i).grupo)"), "Ctrl+Shift+G desagrupa")
    E("() => { SD.sel = [sdSlide().camadas[0].id]; sdAvisar('sel'); }")
    K("Control+]"); ok(E("sdSlide().camadas[1].id === SD.sel[0]"), "Ctrl+] traz para a frente")
    K("Control+["); ok(E("sdSlide().camadas[0].id === SD.sel[0]"), "Ctrl+[ envia para trás")
    K("Control+Shift+]"); ok(E("sdSlide().camadas.at(-1).id === SD.sel[0]"), "Ctrl+Shift+] topo")
    K("Control+l"); ok(E("!!sdCamada(SD.sel[0]).trava"), "Ctrl+L trava"); K("Control+l")
    K("Control+h"); ok(E("!!sdCamada(SD.sel[0]).oculto"), "Ctrl+H oculta"); K("Control+h")
    K("Control+b"); ok(E("+sdCamada(SD.sel[0]).peso >= 600"), "Ctrl+B negrito na camada")
    K("Control+Shift+c"); E("() => { SD.sel = [sdSlide().camadas[0].id]; sdAvisar('sel'); }"); K("Control+Shift+v")
    ok(E("+sdCamada(SD.sel[0]).peso >= 600"), "Ctrl+Shift+C / Ctrl+Shift+V copia e cola o estilo")
    K("Control+c"); K("Control+v"); ok(E("sdSlide().camadas.length") == 3, "Ctrl+C / Ctrl+V")
    K("Control+x"); ok(E("sdSlide().camadas.length") == 2, "Ctrl+X recorta")
    K("Control+z"); ok(E("sdSlide().camadas.length") == 3, "Ctrl+Z desfaz")
    K("Control+y"); ok(E("sdSlide().camadas.length") == 2, "Ctrl+Y refaz")
    E("() => { SD.sel = [sdSlide().camadas[0].id]; sdAvisar('sel'); }")
    K("Enter"); ok(E("!!SDU.editando"), "Enter edita o texto"); K("Escape"); ok(E("!SDU.editando"), "Esc sai da edição")
    K("Escape"); ok(E("SD.sel.length") == 0, "Esc tira a seleção")
    K("Control+1"); ok(abs(E("SDU.z") - 1) < 0.01, "Ctrl+1 = 100%")
    K("Control+="); ok(E("SDU.z") > 1.1, "Ctrl+= amplia")
    z1 = E("SDU.z"); K("Control+-"); ok(E("SDU.z") < z1, "Ctrl+− reduz"); K("Control+0"); ok(E("SDU.ajustar === true"), "Ctrl+0 ajusta à tela")
    i0 = E("SD.atual"); K("PageUp"); ok(E("SD.atual") == i0 - 1, "PageUp volta um slide")
    K("PageDown"); ok(E("SD.atual") == i0, "PageDown avança")
    K("Home"); ok(E("SD.atual") == 0, "Home vai ao primeiro"); K("End"); ok(E("SD.atual") == E("SD.doc.slides.length - 1"), "End vai ao último")
    K("ArrowUp"); ok(E("SD.atual") == E("SD.doc.slides.length - 2"), "seta sem seleção troca de slide")
    pg.click("#sd-lista .sd-it.on .sd-mini"); time.sleep(0.2)
    sid = E("sdSlide().id"); K("Control+ArrowUp"); ok(E("(id) => sdSlide().id === id && SD.atual === SD.doc.slides.length - 3", sid), "Ctrl+↑ na lista muda o slide de lugar")
    n1 = E("SD.doc.slides.length"); K("Delete"); ok(E("SD.doc.slides.length") == n1 - 1, "Delete com a lista em foco apaga o slide")
    K("Control+z"); ok(E("SD.doc.slides.length") == n1, "e Ctrl+Z traz de volta")
    K("F5"); time.sleep(0.6); ok(E("SDS.on"), "F5 apresenta (sem recarregar a página)")
    K("Escape"); time.sleep(0.3); K("Shift+F5"); time.sleep(0.6); ok(E("SDS.on && SDS.lista[SDS.i] === sdSlide()"), "Shift+F5 apresenta deste slide")
    K("b"); ok(E("!SDS.el.querySelector('.sd-show-tela').hidden"), "B tela preta"); K("b")
    K("Escape"); time.sleep(0.3)
    ok(not erros, "sem erro de JS", "; ".join(erros[:3]))

print("\nRESULTADO:", "OK" if not falhas else f"{len(falhas)} FALHA(S): " + "; ".join(falhas))
sys.exit(1 if falhas else 0)
