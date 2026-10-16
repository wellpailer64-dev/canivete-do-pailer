"""Slider Kanivete (Sd): comandos, desfazer, HTML → slide, mídias (imagem, GIF, vídeo), salvar/abrir .dknv, exportar
PDF/PNG/PPTX (editável e imagens)/MP4, reimportar o PPTX, mouse de verdade na mesa (mover, alça, laço, texto) e a
apresentação (passos, transição, sair). App de teste em --agente=9333.
    py -3.13 testes/teste_slider.py [--porta 9333] [--sem-mp4]
Saídas em %TEMP%/kanivete_teste_slider. Sai com 1 se algo reprovar."""
import argparse, json, os, subprocess, sys, tempfile, time
from playwright.sync_api import sync_playwright
sys.stdout.reconfigure(encoding="utf-8")
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
ap = argparse.ArgumentParser(); ap.add_argument("--porta", type=int, default=9333); ap.add_argument("--sem-mp4", action="store_true")
a = ap.parse_args()
OUT = os.path.join(tempfile.gettempdir(), "kanivete_teste_slider").replace("\\", "/"); os.makedirs(OUT, exist_ok=True)
falhas = []


def ok(cond, nome, extra=""):
    print(("  ok  " if cond else "  FALHOU ") + nome + (f"  ({extra})" if extra else ""), flush=True)
    if not cond:
        falhas.append(nome)


# material: PNG, GIF animado e MP4 curtos
from Functions.midia import ffmpeg
FF = ffmpeg()
def gerar(arq, args):
    p = f"{OUT}/{arq}"
    if not os.path.isfile(p):
        subprocess.run([FF, "-y", "-loglevel", "error", *args, p], check=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    return p
PNG = gerar("foto.png", ["-f", "lavfi", "-i", "testsrc2=s=1280x720", "-frames:v", "1"])
GIF = gerar("anim.gif", ["-f", "lavfi", "-i", "testsrc=s=320x240:r=10:d=2"])
MP4 = gerar("clipe.mp4", ["-f", "lavfi", "-i", "testsrc2=s=1280x720:r=30:d=3", "-f", "lavfi", "-i", "sine=f=440:d=3", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest"])

with sync_playwright() as p:
    pg = next(x for c in p.chromium.connect_over_cdp(f"http://127.0.0.1:{a.porta}").contexts for x in c.pages if "index.html" in x.url)
    pg.set_default_timeout(120000)
    erros = []
    pg.on("pageerror", lambda e: erros.append(str(e)))
    E = lambda js, arg=None: pg.evaluate(js, arg)
    E("switchTool('slider-kanivete')"); time.sleep(1)
    ok(E("!document.getElementById('sd-inicio').hidden"), "lobby aparece ao entrar sem documento aberto") if not E("!!SD.doc") else None

    print("comandos")
    E("async () => { await SDK.cmd('modelo', {tema: 'escuro'}); SDK_UI.editor(); }"); time.sleep(0.8)
    m = E("SDK.mapa()")
    ok(len(m["slides"]) == 8, "modelo escuro com 8 slides")
    ok(E("sdLum(SD.doc.slides[0].fundo.cor) < 0.2"), "capa do tema escuro é escura")
    r = E("""async ([png, gif, mp4]) => {
        const out = {};
        out.t = await SDK.cmd('texto', {texto: 'Olá, plateia', x: 100, y: 100, w: 800, h: 120, tam: 72, slide: 3});
        out.f = await SDK.cmd('forma', {forma: 'estrela', x: 1200, y: 200, w: 300, h: 300, slide: 3});
        out.i = await SDK.cmd('imagem', {src: png, slide: 3});
        out.g = await SDK.cmd('imagem', {src: gif, x: 100, y: 600, w: 320, h: 240, slide: 3});
        out.v = await SDK.cmd('video', {src: mp4, x: 1300, y: 600, w: 480, h: 270, slide: 3, loop: true, mudo: true});
        await SDK.cmd('slide_ir', {slide: 3});
        return out; }""", [PNG, GIF, MP4])
    ids = {k: v["id"] for k, v in r.items()}
    SM = E("sdSlide(3).id")   # o slide das mídias (seguido pelo id: mover slides muda o número)
    s3 = E("SDK.mapa().slides[2].camadas.map(c => c.tipo)")
    ok(s3.count("imagem") == 2 and "video" in s3 and "forma" in s3, "texto, forma, imagem, GIF e vídeo no slide 3", ",".join(s3))
    img = E("(id) => { const c = sdCamada(id); return [c.w, c.h]; }", ids["i"])
    ok(abs(img[0] / img[1] - 16 / 9) < 0.02, "imagem entra na proporção do arquivo", str(img))
    n0 = E("SD.hist.length")
    E("(id) => SDK.cmd('editar', {ids: [id], cor: '#ff0000', tam: 90})", ids["t"])
    ok(E("(id) => sdCamada(id).cor === '#ff0000' && sdCamada(id).tam === 90", ids["t"]), "editar propriedades")
    E("(ids) => SDK.cmd('alinhar', {ids, como: 'esq', a: 'selecao'})", [ids["t"], ids["f"]])
    ok(E("([a, b]) => sdCamada(a).x === sdCamada(b).x", [ids["t"], ids["f"]]), "alinhar à esquerda da seleção")
    E("(ids) => SDK.cmd('distribuir', {ids, eixo: 'h'})", [ids["t"], ids["f"], ids["g"]])
    E("(id) => SDK.cmd('ordem', {ids: [id], para: 'fundo'})", ids["f"])
    ok(E("(id) => sdSlide().camadas[0].id === id", ids["f"]), "enviar para o fundo")
    d = E("(id) => SDK.cmd('duplicar', {ids: [id]})", ids["t"])
    ok(len(d["ids"]) == 1 and E("sdSlide().camadas.length") == len(s3) + 1, "duplicar camada")
    E("(id) => SDK.cmd('apagar', {ids: [id]})", d["ids"][0])
    E("(id) => SDK.cmd('animacao', {ids: [id], tipo: 'subir', clique: true})", ids["t"])
    ok(E("SD.hist.length") > n0, "histórico registra as mudanças")
    E("SDK.desfazer()")
    ok(E("(id) => !sdCamada(id).anim", ids["t"]), "desfazer volta a animação")
    E("SDK.refazer()")
    ok(E("(id) => sdCamada(id).anim && sdCamada(id).anim.tipo === 'subir'", ids["t"]), "refazer")

    r = E("SDK.cmd('tabela', {linhas: 3, colunas: 3, textos: [['A', 'B', 'C'], ['1', '2', '3']]})")
    ok(len(r["ids"]) == 18 and E("(g) => sdSlide().camadas.filter(c => c.grupo === g).length", r["grupo"]) == 18, "tabela 3×3 agrupada (formas + textos)")
    E("(ids) => SDK.cmd('desagrupar', {ids})", r["ids"][:1])
    ok(E("sdSlide().camadas.every(c => !c.grupo)"), "desagrupar")
    g = E("(ids) => SDK.cmd('agrupar', {ids})", r["ids"][:4])
    ok(E("(g) => sdSlide().camadas.filter(c => c.grupo === g).length", g["grupo"]) == 4, "agrupar")
    E("(ids) => SDK.cmd('apagar', {ids})", r["ids"])

    print("slides")
    E("SDK.cmd('slide_novo', {layout: 'lista'})"); E("SDK.cmd('slide_duplicar', {})")
    ok(len(E("SDK.mapa()")["slides"]) == 10, "novo e duplicar slide")
    lis = E("SDK.mapa().slides.filter(s => s.camadas.some(c => c.nome === 'Tópico 1')).map(s => s.n)")
    E("(n) => SDK.cmd('slide_mover', {slide: n, para: 1})", lis[-1])
    ok(E("SDK.mapa().slides[0].camadas.some(c => c.nome === 'Tópico 1')"), "mover slide para o começo")
    E("SDK.cmd('slide_apagar', {slide: 1})")
    E("(n) => SDK.cmd('slide_apagar', {slide: n})", E("SDK.mapa().slides.filter(s => s.camadas.some(c => c.nome === 'Tópico 1')).map(s => s.n)")[0])
    ok(len(E("SDK.mapa()")["slides"]) == 8, "apagar slides")
    E("SDK.cmd('slide_fundo', {slide: 2, tipo: 'degrade', cor: '#0e1116', cor2: '#145a52', angulo: 120})")
    E("SDK.cmd('slide_transicao', {tipo: 'deslizar', dur: 0.4, todos: true})")
    E("(id) => SDK.cmd('slide_props', {slide: id, notas: 'Lembrar de mostrar o vídeo.'})", SM)

    print("HTML → slide")
    r = E("""SDK.cena(`<div style="width:1920px;height:1080px;background:#101820;position:relative;font-family:Jost">
        <h1 style="position:absolute;left:140px;top:160px;margin:0;font-size:120px;font-weight:600;color:#fff">Resultado do mês</h1>
        <div style="position:absolute;left:140px;top:420px;width:520px;height:300px;border-radius:24px;background:#2ee6c8"></div>
        <p style="position:absolute;left:760px;top:440px;width:900px;margin:0;font-size:48px;color:#c8d3dc">Texto <b>com destaque</b> e quebra<br>de linha.</p>
      </div>`, {novo: true})""")
    m = E("SDK.mapa()")
    sl = m["slides"][r["slide"] - 1]
    ok(sl["fundo"].get("cor") == "#101820" and len(sl["camadas"]) == 3, "HTML vira fundo + 3 camadas", json.dumps([c["tipo"] for c in sl["camadas"]]))
    ok(E("(n) => sdSlide(n).camadas.find(c => c.tipo === 'texto' && /destaque/.test(c.html)).html.includes('<b>')", r["slide"]), "negrito do HTML vira trecho")

    print("fontes por arquivo")
    E("async () => { await SDK.cmd('slide_ir', {slide: 1}); await SDK.cmd('texto', {texto: 'Fraunces aqui', fonte: 'Fraunces', tam: 90, x: 100, y: 850, w: 900, h: 120}); }")
    time.sleep(2)
    ok(E("!sdFonteNativa('Fraunces') ? document.fonts.check('40px Fraunces') && !!SD.fontesArq['Fraunces'] : true"), "família que o navegador não vê entra por arquivo (Fraunces)")
    ok(E("async () => (await sdCssFontes(sdSlide(1))).length > 1000 || sdFonteNativa('Fraunces')"), "a foto de exportação leva a fonte embutida")

    print("foto do slide")
    foto = E("async (SM) => { const d = await sdFotoSlide(sdSlide(SM), 480, 'image/png'); const im = new Image(); im.src = d; await im.decode();"
             " const cv = document.createElement('canvas'); cv.width = 480; cv.height = 270; const g = cv.getContext('2d'); g.drawImage(im, 0, 0);"
             " const px = g.getImageData(0, 0, 480, 270).data; let cores = new Set(); for (let i = 0; i < px.length; i += 4 * 97) cores.add(px[i] >> 5 << 10 | px[i + 1] >> 5 << 5 | px[i + 2] >> 5);"
             " return {n: cores.size, len: d.length}; }", SM)
    ok(foto["n"] > 20, "foto do slide 3 tem a imagem/vídeo (muitas cores)", str(foto))

    print("salvar e abrir")
    dk = f"{OUT}/palestra.dknv"
    if os.path.exists(dk): os.remove(dk)
    r = E("(c) => SDK.cmd('salvar', {caminho: c})", dk)
    ok(os.path.isfile(dk) and os.path.getsize(dk) > 50000, ".dknv salvo com as mídias dentro", f"{os.path.getsize(dk) // 1024} KB" if os.path.isfile(dk) else "")
    antes = E("JSON.stringify(SDK.mapa().slides.map(s => s.camadas.length))")
    E("SDK.cmd('novo', {})")
    E("(c) => SDK.cmd('abrir', {caminho: c})", dk)
    ok(E("JSON.stringify(SDK.mapa().slides.map(s => s.camadas.length))") == antes, "reabrir o .dknv traz os mesmos slides/camadas")
    ok(E("(id) => sdSlide(id).camadas.filter(c => c.src).length === 3 && sdSlide(id).camadas.filter(c => c.src).every(c => /^[A-Z]:/i.test(c.src))", SM), "mídias voltam com caminho absoluto")

    print("exportar")
    N = len(E("SDK.mapa()")["slides"])
    for fmt, arq in (("pdf", "palestra.pdf"), ("pptx", "palestra.pptx"), ("pptx_imagens", "palestra_img.pptx"), ("png", "pngs")):
        alvo = f"{OUT}/{arq}"
        t0 = time.time(); r = E("([f, c]) => SDK.cmd('exportar', {formato: f, caminho: c, largura: 1280})", [fmt, alvo])
        existe = os.path.isdir(alvo) and len(os.listdir(alvo)) == N if fmt == "png" else os.path.isfile(alvo) and os.path.getsize(alvo) > 10000
        ok(existe, f"exportar {fmt}", f"{time.time() - t0:.1f}s")
    from pptx import Presentation
    prs = Presentation(f"{OUT}/palestra.pptx")
    tipos = [sh.shape_type for s in prs.slides for sh in s.shapes]
    ok(len(prs.slides) == N and any(sh.has_text_frame and "Resultado do mês" in sh.text_frame.text for s in prs.slides for sh in s.shapes), f"PPTX editável: {N} slides com texto de verdade")
    ok(any(str(t).startswith("MEDIA") or "MEDIA" in str(t) for t in tipos), "PPTX editável leva o vídeo como vídeo", str(set(map(str, tipos)))[:120])
    ok(any(sl.has_notes_slide and "vídeo" in sl.notes_slide.notes_text_frame.text for sl in prs.slides), "notas do apresentador no PPTX")
    r = E("(c) => SDK.cmd('abrir', {caminho: c})", f"{OUT}/palestra.pptx")
    ok(r and r.get("slides") == N and E("SDK.mapa().slides.some(s => s.camadas.some(c => c.texto && c.texto.includes('Resultado')))"), "reimportar o PPTX exportado")
    E("(c) => SDK.cmd('abrir', {caminho: c})", dk)
    if not a.sem_mp4:
        t0 = time.time(); E("(c) => SDK.cmd('exportar', {formato: 'mp4', caminho: c, largura: 1280})", f"{OUT}/palestra.mp4")
        dur = subprocess.run([FF.replace("ffmpeg.exe", "ffprobe.exe"), "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", f"{OUT}/palestra.mp4"], capture_output=True, text=True).stdout.strip() if os.path.isfile(FF.replace("ffmpeg.exe", "ffprobe.exe")) else "?"
        ok(os.path.isfile(f"{OUT}/palestra.mp4") and os.path.getsize(f"{OUT}/palestra.mp4") > 50000, "exportar MP4 com transições", f"{time.time() - t0:.1f}s, {dur}s de vídeo")

    print("mouse na mesa")
    E("async () => { await SDK.cmd('slide_novo', {layout: 'vazio'}); await SDK.cmd('forma', {forma: 'ret', x: 1300, y: 700, w: 300, h: 200});"
      " await SDK.cmd('forma', {forma: 'elipse', x: 1650, y: 700, w: 200, h: 200}); SD.sel = []; sdAvisar('sel'); }"); time.sleep(0.4)
    tid = E("(async () => (await SDK.cmd('texto', {texto: 'Arraste', x: 300, y: 300, w: 600, h: 120, tam: 80})).id)()")
    E("SD.sel = []; sdAvisar('sel')"); time.sleep(0.4)
    caixa = lambda i: pg.locator(f'#sd-palco [data-id="{i}"]').bounding_box()
    b = caixa(tid); x0 = E("(i) => sdCamada(i).x", tid)
    pg.mouse.move(b["x"] + 20, b["y"] + 10); pg.mouse.down(); pg.mouse.move(b["x"] + 80, b["y"] + 60, steps=6); pg.mouse.up(); time.sleep(0.3)
    ok(E("(i) => sdCamada(i).x", tid) > x0 + 30, "arrastar move a camada", f"{x0} → {E('(i) => sdCamada(i).x', tid)}")
    w0 = E("(i) => sdCamada(i).w", tid)
    al = pg.locator('.sd-alca.a-e').bounding_box()
    pg.mouse.move(al["x"] + 5, al["y"] + 5); pg.mouse.down(); pg.mouse.move(al["x"] + 105, al["y"] + 5, steps=6); pg.mouse.up(); time.sleep(0.3)
    ok(E("(i) => sdCamada(i).w", tid) > w0 + 50, "alça da direita alarga", f"{w0} → {E('(i) => sdCamada(i).w', tid)}")
    pb = pg.locator('#sd-papel').bounding_box()
    pg.mouse.move(pb["x"] + 3, pb["y"] + 3); pg.mouse.down(); pg.mouse.move(pb["x"] + pb["width"] - 3, pb["y"] + pb["height"] - 3, steps=8); pg.mouse.up(); time.sleep(0.3)
    ok(E("SD.sel.length") == E("sdSlide().camadas.length") == 3, "laço seleciona as 3 camadas", str(E("SD.sel.length")))
    pg.keyboard.press("Escape"); time.sleep(0.2)
    b = caixa(tid); pg.mouse.dblclick(b["x"] + b["width"] / 2, b["y"] + b["height"] / 2); time.sleep(0.3)
    pg.keyboard.type("Texto digitado"); pg.mouse.click(pb["x"] + 5, pb["y"] + pb["height"] - 5); time.sleep(0.4)
    ok(E("(i) => sdTextoPuro(sdCamada(i)) === 'Texto digitado'", tid), "editar texto no lugar (duplo clique e digitar)", E("(i) => sdTextoPuro(sdCamada(i))", tid))
    n = E("sdSlide().camadas.length"); E("(i) => { SD.sel = [i]; sdAvisar('sel'); }", tid)
    pg.locator('#sd-mesa').focus() if False else None
    pg.keyboard.press("Delete"); time.sleep(0.2)
    ok(E("sdSlide().camadas.length") == n - 1, "Delete apaga a seleção")
    pg.keyboard.press("Control+z"); time.sleep(0.2)
    ok(E("sdSlide().camadas.length") == n, "Ctrl+Z desfaz pelo teclado")
    pg.screenshot(path=f"{OUT}/editor.jpg", quality=70)

    print("apresentação")
    E("(id) => SDK.cmd('slide_ir', {slide: id})", SM)
    E("sdApresentar(SD.atual)"); time.sleep(1.2)
    st = E("SDS_API.estado()")
    ok(st["on"] and st["i"] == E("SD.doc.slides.filter(s => !s.oculto).indexOf(sdSlide())"), "apresentação abre no slide escolhido", str(st))
    ok(E("!!document.querySelector('.sd-show video') && !document.querySelector('.sd-show video').paused"), "vídeo toca sozinho no slide")
    pg.screenshot(path=f"{OUT}/apresentacao.jpg", quality=70)
    pg.keyboard.press("ArrowRight"); time.sleep(0.6)
    st2 = E("SDS_API.estado()")
    ok(st2["i"] == 3 or st2["passo"] >= 0, "seta avança", str(st2))
    pg.keyboard.press("n"); time.sleep(0.4)
    ok(E("!document.querySelector('.sd-show-ap').hidden"), "N abre a visão do apresentador")
    pg.screenshot(path=f"{OUT}/apresentador.jpg", quality=70)
    pg.keyboard.press("Escape"); time.sleep(0.6)
    ok(not E("SDS_API.estado().on") and not E("!!document.querySelector('.sd-show')"), "Esc sai da apresentação")

    ok(not erros, "sem erro de JS", "; ".join(erros[:4]))
print("\nRESULTADO:", "OK" if not falhas else f"{len(falhas)} falha(s): " + "; ".join(falhas))
sys.exit(1 if falhas else 0)
