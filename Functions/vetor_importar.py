"""Vetor Kanivete: abrir PDF, AI, SVG e PowerPoint como documento vetorial editável (modelo em Instructions/vetor-kanivete.md).

- PDF / AI (o .ai salvo "compatível com PDF", o padrão do Illustrator): interpretador do content stream (pikepdf) —
  caminhos, cores EXATAS (CMYK, cinza, RGB, Separation = cor especial, registro), traço, tracejado, opacidade, mesclagem,
  sobreimpressão, máscaras de corte (W n), degradês (sh / padrão de sombreamento tipo 2 e 3), imagens (CMYK fica CMYK),
  Form XObjects e as CAMADAS do Illustrator (conteúdo marcado /OC). Texto vem do PyMuPDF (fonte, tamanho, posição) com a
  cor tirada do próprio content stream (o PyMuPDF só dá RGB).
- SVG: caminhos (inclusive arcos), formas, estilos (atributo, style e as classes .cls-N que o Illustrator exporta),
  transformações, degradês, clipPath, use, texto e imagens.
- PPTX: um slide = uma prancheta; formas (preset e livre), preenchimento/linha (cores do tema), textos, imagens, grupos.
O que não vira objeto editável entra no `relatorio` (o painel mostra ao abrir)."""
import base64, hashlib, io, math, os, re, tempfile, zlib
from collections import Counter

PASTA = os.path.join(tempfile.gettempdir(), "vetor_kanivete", "importados")
GAP = 36


class _Ids:
    def __init__(self):
        self.n = 0

    def __call__(self, p="o"):
        self.n += 1
        return f"{p}{self.n}"


def _doc(nome, modo="cmyk"):
    return {"versao": 1, "nome": nome, "unidade": "mm", "modoCor": modo, "perfil": "FOGRA39", "sangria": round(3 * 72 / 25.4, 4),
            "pranchetas": [], "camadas": [], "amostras": [], "imagens": {}, "guias": []}


def _pt(x, y):
    return [x, y, x, y, x, y]


# ─────────────────────────── matrizes (convenção PDF [a b c d e f]) ───────────────────────────
def mmul(m1, m2):
    a1, b1, c1, d1, e1, f1 = m1
    a2, b2, c2, d2, e2, f2 = m2
    return [a1 * a2 + b1 * c2, a1 * b2 + b1 * d2, c1 * a2 + d1 * c2, c1 * b2 + d1 * d2, e1 * a2 + f1 * c2 + e2, e1 * b2 + f1 * d2 + f2]


def mpt(m, x, y):
    return m[0] * x + m[2] * y + m[4], m[1] * x + m[3] * y + m[5]


def mescala(m):
    return math.sqrt(abs(m[0] * m[3] - m[1] * m[2])) or 1


def minv(m):
    a, b, c, d, e, f = m
    det = a * d - b * c or 1e-12
    return [d / det, -b / det, -c / det, a / det, (c * f - d * e) / det, (b * e - a * f) / det]


def _transformar_subs(subs, m):
    for sb in subs:
        for p in sb["pts"]:
            for i in (0, 2, 4):
                p[i], p[i + 1] = mpt(m, p[i], p[i + 1])
    return subs


def _ret_subs(x, y, w, h, r=0):
    if r <= 0:
        return [{"fechado": True, "pts": [_pt(x, y), _pt(x + w, y), _pt(x + w, y + h), _pt(x, y + h)]}]
    r = min(r, w / 2, h / 2); k = r * 0.5523
    return [{"fechado": True, "pts": [
        [x + r, y, x + r - k, y, x + r, y], [x + w - r, y, x + w - r, y, x + w - r + k, y],
        [x + w, y + r, x + w, y + r - k, x + w, y + r], [x + w, y + h - r, x + w, y + h - r, x + w, y + h - r + k],
        [x + w - r, y + h, x + w - r + k, y + h, x + w - r, y + h], [x + r, y + h, x + r, y + h, x + r - k, y + h],
        [x, y + h - r, x, y + h - r + k, x, y + h - r], [x, y + r, x, y + r, x, y + r - k]]}]


def _elipse_subs(cx, cy, rx, ry):
    k = 0.5523
    return [{"fechado": True, "pts": [
        [cx + rx, cy, cx + rx, cy - ry * k, cx + rx, cy + ry * k], [cx, cy + ry, cx + rx * k, cy + ry, cx - rx * k, cy + ry],
        [cx - rx, cy, cx - rx, cy + ry * k, cx - rx, cy - ry * k], [cx, cy - ry, cx - rx * k, cy - ry, cx + rx * k, cy - ry]]}]


def _bbox_subs(subs):
    xs = [p[0] for s in subs for p in s["pts"]]; ys = [p[1] for s in subs for p in s["pts"]]
    return (min(xs), min(ys), max(xs), max(ys)) if xs else (0, 0, 0, 0)


def _salvar_imagem(im, doc, ids, nome_base="img", mascara=None):
    """PIL → arquivo em PASTA (CMYK vira TIFF CMYK, o resto PNG) e registro em doc.imagens. → id
    mascara (PIL L) = transparência de imagem CMYK (fica num PNG ao lado: o TIFF continua CMYK para a gráfica)."""
    os.makedirs(PASTA, exist_ok=True)
    iid = ids("i")
    h = hashlib.md5(im.tobytes() + im.mode.encode() + str(im.size).encode()).hexdigest()[:16]   # inteiro: topo igual (branco) não pode colidir
    extra = {}
    if im.mode == "CMYK":
        arq = os.path.join(PASTA, f"{nome_base}_{h}.tif"); im.save(arq, compression="tiff_lzw")
        if mascara is not None and mascara.getextrema() != (255, 255):
            if mascara.size != im.size: mascara = mascara.resize(im.size)
            ma = os.path.join(PASTA, f"{nome_base}_{h}_mascara.png"); mascara.save(ma); extra = {"mascara": ma, "alfa": True}
    else:
        if im.mode not in ("RGB", "RGBA", "L", "LA"):
            im = im.convert("RGBA" if "A" in im.mode or im.mode == "P" else "RGB")
        arq = os.path.join(PASTA, f"{nome_base}_{h}.png"); im.save(arq)
    doc["imagens"][iid] = {"arquivo": arq, "w": im.width, "h": im.height, "modo": im.mode, "nome": os.path.basename(arq), **extra}
    return iid


# ═══════════════════════════ PDF / AI ═══════════════════════════
def _func_eval(fn, x):
    """Avalia uma função PDF (tipo 0 8 bits, 2, 3, 4) em x → lista de saídas."""
    t = int(fn.get("/FunctionType", 2))
    dom = [float(v) for v in fn.get("/Domain", [0, 1])]
    x = min(max(x, dom[0]), dom[1])
    if t == 2:
        c0 = [float(v) for v in fn.get("/C0", [0])]; c1 = [float(v) for v in fn.get("/C1", [1])]; n = float(fn.get("/N", 1))
        return [a + (x ** n) * (b - a) for a, b in zip(c0, c1)]
    if t == 3:
        fns, bounds, enc = fn["/Functions"], [float(v) for v in fn.get("/Bounds", [])], [float(v) for v in fn["/Encode"]]
        lim = [dom[0]] + bounds + [dom[1]]
        k = next((i for i in range(len(bounds)) if x < bounds[i]), len(bounds))
        a, b = lim[k], lim[k + 1]; e0, e1 = enc[2 * k], enc[2 * k + 1]
        return _func_eval(fns[k], e0 + (x - a) * (e1 - e0) / ((b - a) or 1))
    if t == 0:
        try:
            dados = fn.read_bytes(); n_out = len(fn["/Range"]) // 2; size = int(fn["/Size"][0])
            if int(fn.get("/BitsPerSample", 8)) != 8:
                raise ValueError
            i = round((x - dom[0]) / ((dom[1] - dom[0]) or 1) * (size - 1))
            rng = [float(v) for v in fn["/Range"]]
            return [rng[2 * j] + dados[i * n_out + j] / 255 * (rng[2 * j + 1] - rng[2 * j]) for j in range(n_out)]
        except Exception:
            return [x]
    if t == 4:
        try:
            src = fn.read_bytes().decode("latin1")
            return _ps_calc(src, [x])
        except Exception:
            return [x]
    return [x]


def _ps_calc(src, pilha):
    toks = re.findall(r"[{}]|[^\s{}]+", src)

    def bloco(i):
        corpo, prof = [], 0
        while i < len(toks):
            t = toks[i]
            if t == "{":
                if prof == 0 and corpo == [] and False:
                    pass
                sub, i = bloco(i + 1); corpo.append(sub); continue
            if t == "}":
                return corpo, i + 1
            corpo.append(t); i += 1
        return corpo, i

    prog, _ = bloco(1 if toks and toks[0] == "{" else 0)

    def rodar(prog, s):
        for t in prog:
            if isinstance(t, list):
                s.append(t); continue
            if re.fullmatch(r"-?[\d.]+(e-?\d+)?", t):
                s.append(float(t)); continue
            if t in ("add", "sub", "mul", "div", "exp", "max", "min"):
                b, a = s.pop(), s.pop()
                s.append({"add": a + b, "sub": a - b, "mul": a * b, "div": a / b if b else 0, "exp": a ** b, "max": max(a, b), "min": min(a, b)}[t])
            elif t in ("gt", "lt", "ge", "le", "eq", "ne"):
                b, a = s.pop(), s.pop(); s.append({"gt": a > b, "lt": a < b, "ge": a >= b, "le": a <= b, "eq": a == b, "ne": a != b}[t])
            elif t == "neg": s.append(-s.pop())
            elif t == "abs": s.append(abs(s.pop()))
            elif t in ("cvr", "cvi", "truncate", "floor", "round", "ceiling"): s.append(float(int(s.pop())) if t != "cvr" else s.pop())
            elif t == "dup": s.append(s[-1])
            elif t == "exch": s[-1], s[-2] = s[-2], s[-1]
            elif t == "pop": s.pop()
            elif t == "copy": n = int(s.pop()); s.extend(s[-n:] if n else [])
            elif t == "index": n = int(s.pop()); s.append(s[-1 - n])
            elif t == "roll":
                j, n = int(s.pop()), int(s.pop())
                if n:
                    seg = s[-n:]; j %= n; s[-n:] = seg[-j:] + seg[:-j]
            elif t == "if":
                p = s.pop(); c = s.pop()
                if c: rodar(p, s)
            elif t == "ifelse":
                p2, p1, c = s.pop(), s.pop(), s.pop(); rodar(p1 if c else p2, s)
            elif t == "true": s.append(True)
            elif t == "false": s.append(False)
        return s
    return [float(v) for v in rodar(prog, list(pilha))]


class _PDF:
    def __init__(self, pdf, doc, ids, rel):
        self.pdf, self.doc, self.ids, self.rel = pdf, doc, ids, rel
        self.camadas = {}
        self.textos = []   # (x, y, cor) de cada operador de mostrar texto
        self.cmyk = 0

    # cores
    def _cs(self, cs, res):
        import pikepdf
        if isinstance(cs, pikepdf.Name):
            n = str(cs)
            if n in ("/DeviceGray", "/G", "/CalGray"): return ("gray", None)
            if n in ("/DeviceRGB", "/RGB", "/CalRGB"): return ("rgb", None)
            if n in ("/DeviceCMYK", "/CMYK"): return ("cmyk", None)
            if n == "/Pattern": return ("pattern", None)
            r = res.get("/ColorSpace", {}) if res is not None else {}
            if n in r: return self._cs(r[n], res)
            return ("gray", None)
        if isinstance(cs, pikepdf.Array) and len(cs):
            t = str(cs[0])
            if t == "/ICCBased":
                n = int(cs[1].get("/N", 3)); return {1: ("gray", None), 3: ("rgb", None), 4: ("cmyk", None)}.get(n, ("rgb", None))
            if t == "/Separation":
                return ("sep", (str(cs[1])[1:], cs[2], cs[3]))
            if t == "/DeviceN":
                return ("devn", ([str(x)[1:] for x in cs[1]], cs[2], cs[3]))
            if t == "/Indexed":
                return ("idx", (self._cs(cs[1], res), int(cs[2]), cs[3]))
            if t == "/Pattern":
                return ("pattern", None)
            if t in ("/CalRGB", "/Lab"): return ("rgb", None) if t == "/CalRGB" else ("lab", None)
            if t == "/CalGray": return ("gray", None)
        return ("gray", None)

    def _alt_cmyk(self, alt, fn, ent, res):
        try:
            sai = _func_eval(fn, ent[0]) if len(ent) == 1 else _func_eval(fn, ent[0])
        except Exception:
            sai = [0, 0, 0, 1]
        cs = self._cs(alt, res)
        c = self._cor(cs, sai, res)
        if c and c["k"] == "cmyk": return c["v"]
        if c and c["k"] == "rgb":
            from Functions.vetor_kanivete import rgb_para_cmyk
            return rgb_para_cmyk([c["v"]])[0]
        return [0, 0, 0, 100]

    def _cor(self, cs, comps, res):
        tipo, info = cs
        comps = [float(v) for v in comps if not hasattr(v, "is_name")] if comps else []
        try:
            if tipo == "gray":
                return {"k": "cmyk", "v": [0, 0, 0, round((1 - (comps[0] if comps else 0)) * 100, 2)]}
            if tipo == "rgb":
                c = (comps + [0, 0, 0])[:3]; return {"k": "rgb", "v": [round(x * 255) for x in c]}
            if tipo == "cmyk":
                self.cmyk += 1
                c = (comps + [0, 0, 0, 0])[:4]; return {"k": "cmyk", "v": [round(x * 100, 2) for x in c]}
            if tipo == "sep":
                nome, alt, fn = info
                t = comps[0] if comps else 1
                if nome == "All": return {"k": "reg"}
                if nome == "None": return None
                self.cmyk += 1
                return {"k": "spot", "nome": nome, "v": [round(x, 2) for x in self._alt_cmyk(alt, fn, [1.0], res)], "tint": round(t * 100, 1)}
            if tipo == "devn":
                nomes, alt, fn = info
                proc = {"Cyan": 0, "Magenta": 1, "Yellow": 2, "Black": 3}
                if all(n in proc for n in nomes):
                    v = [0, 0, 0, 0]
                    for n, x in zip(nomes, comps): v[proc[n]] = round(x * 100, 2)
                    self.cmyk += 1
                    return {"k": "cmyk", "v": v}
                sai = _func_eval(fn, comps[0]) if fn is not None and len(comps) == 1 else None
                self.rel["devn"] += 1
                return self._cor(self._cs(alt, res), sai or [0, 0, 0, 1], res)
            if tipo == "idx":
                base, hival, tab = info
                dados = tab.read_bytes() if hasattr(tab, "read_bytes") else bytes(tab)
                n = {"gray": 1, "rgb": 3, "cmyk": 4}.get(base[0], 3); i = int(comps[0]) if comps else 0
                return self._cor(base, [b / 255 for b in dados[i * n:(i + 1) * n]], res)
            if tipo == "lab":
                L, a, b = (comps + [0, 0, 0])[:3]
                return {"k": "rgb", "v": [max(0, min(255, round(L * 2.55 + a)))] * 1 + [max(0, min(255, round(L * 2.55)))] + [max(0, min(255, round(L * 2.55 - b)))]}
        except Exception:
            pass
        return {"k": "cmyk", "v": [0, 0, 0, 100]}

    def _gradiente(self, sh, m, res):
        """Shading tipo 2/3 → cor {k:'grad'} em coordenadas do documento (m = matriz do espaço do sombreamento → doc)."""
        t = int(sh.get("/ShadingType", 0))
        if t not in (2, 3):
            self.rel["sombreamento"] += 1
            return None
        cs = self._cs(sh["/ColorSpace"], res); fn = sh["/Function"]
        dom = [float(v) for v in sh.get("/Domain", [0, 1])]
        fns = list(fn) if hasattr(fn, "__len__") and not hasattr(fn, "keys") else None
        def ev(x):
            return [_func_eval(f, x)[0] for f in fns] if fns else _func_eval(fn, x)
        # paradas: limites das funções costuradas (tipo 3) ou 2 pontas; tipo 2 com N≠1 ganha pontos no meio
        xs = [0.0, 1.0]
        f0 = fns[0] if fns else fn
        if int(f0.get("/FunctionType", 2)) == 3:
            xs = [0.0] + [float(b) for b in f0.get("/Bounds", [])] + [1.0]
        elif int(f0.get("/FunctionType", 2)) in (0, 4) or float(f0.get("/N", 1)) != 1:
            xs = [i / 8 for i in range(9)]
        paradas = [{"p": round(x, 4), "cor": self._cor(cs, ev(dom[0] + x * (dom[1] - dom[0])), res)} for x in xs]
        co = [float(v) for v in sh["/Coords"]]
        if t == 2:
            a, b = mpt(m, co[0], co[1]), mpt(m, co[2], co[3])
            return {"k": "grad", "tipo": "lin", "a": [round(a[0], 3), round(a[1], 3)], "b": [round(b[0], 3), round(b[1], 3)], "paradas": paradas}
        c = mpt(m, co[3], co[4]); f = mpt(m, co[0], co[1])
        return {"k": "grad", "tipo": "rad", "a": [round(c[0], 3), round(c[1], 3)], "f": [round(f[0], 3), round(f[1], 3)],
                "r": round(co[5] * mescala(m), 3), "paradas": paradas}

    def _camada(self, nome):
        if nome not in self.camadas:
            c = {"id": self.ids("c"), "nome": nome, "visivel": True, "trava": False, "itens": []}
            self.camadas[nome] = c; self.doc["camadas"].append(c)
        return self.camadas[nome]["itens"]

    def pagina(self, page, F):
        gs = {"ctm": F, "fill": {"k": "cmyk", "v": [0, 0, 0, 100]}, "stroke": {"k": "cmyk", "v": [0, 0, 0, 100]},
              "fcs": ("gray", None), "scs": ("gray", None), "lw": 1, "cap": 0, "join": 0, "miter": 10, "dash": ([], 0),
              "ca": 1, "CA": 1, "bm": "normal", "op": False, "OP": False, "cont": self._camada("Camada 1"), "raiz": True, "clip": None}
        self.camada_pilha = ["Camada 1"]
        self._stream(page.obj, page.Resources if "/Resources" in page.obj else None, gs, 0)

    def _stream(self, alvo, res, gs, prof):
        import pikepdf
        if prof > 12:
            return
        try:
            ops = pikepdf.parse_content_stream(alvo)
        except Exception as e:
            self.rel["erros"].append(f"conteúdo ilegível: {e}"); return
        pilha, subs, cur, clip_pend = [], [], None, None
        tm = tlm = [1, 0, 0, 1, 0, 0]; tl = 0
        bm_map = {"/Multiply": "multiplicacao", "/Screen": "divisao", "/Overlay": "sobrepor", "/SoftLight": "luz_suave", "/HardLight": "luz_forte",
                  "/Darken": "escurecer", "/Lighten": "clarear", "/Difference": "diferenca", "/Exclusion": "exclusao", "/ColorDodge": "subexposicao",
                  "/ColorBurn": "superexposicao", "/Hue": "matiz", "/Saturation": "saturacao", "/Color": "cor", "/Luminosity": "luminosidade"}

        def P(x, y):
            return list(mpt(gs["ctm"], x, y))

        def emitir(preench, traco, regra):
            nonlocal subs, cur, clip_pend
            if subs and (preench or traco):
                o = {"id": self.ids(), "tipo": "caminho", "subs": [dict(s, pts=[list(p) for p in s["pts"]]) for s in subs], "regra": regra,
                     "preench": dict(gs["fill"]) if preench and gs["fill"] else None,
                     "traco": ({"cor": gs["stroke"], "larg": round(gs["lw"] * mescala(gs["ctm"]), 4), "cap": ["butt", "round", "square"][gs["cap"]],
                                "junc": ["miter", "round", "bevel"][gs["join"]], "miter": gs["miter"],
                                "tracejado": [round(d * mescala(gs["ctm"]), 3) for d in gs["dash"][0]], "fase": gs["dash"][1]} if traco and gs["stroke"] else None)}
                if preench and gs["ca"] < 1 or traco and gs["CA"] < 1:
                    o["op"] = round(gs["ca"] if preench else gs["CA"], 3)
                if gs["bm"] != "normal": o["bm"] = gs["bm"]
                if (preench and gs["op"]) or (traco and gs["OP"]): o["sobre"] = {"p": bool(preench and gs["op"]), "t": bool(traco and gs["OP"])}
                if o["preench"] or o["traco"]:
                    gs["cont"].append(o)
            if clip_pend and subs:
                g = {"id": self.ids(), "tipo": "grupo", "clip": True,
                     "itens": [{"id": self.ids(), "tipo": "caminho", "subs": [dict(s, pts=[list(p) for p in s["pts"]]) for s in subs], "regra": clip_pend,
                                "preench": None, "traco": None}]}
                gs["cont"].append(g); gs["cont"] = g["itens"]; gs["raiz"] = False; gs["clip"] = g["itens"][0]["subs"]
            subs, cur, clip_pend = [], None, None

        for operandos, op in ops:
            op = str(op)
            try:
                o = operandos
                if op == "q":
                    pilha.append(dict(gs))
                elif op == "Q":
                    if pilha: gs = pilha.pop()
                elif op == "cm":
                    gs["ctm"] = mmul([float(v) for v in o], gs["ctm"])
                elif op == "m":
                    x, y = P(float(o[0]), float(o[1])); cur = {"fechado": False, "pts": [_pt(x, y)]}; subs.append(cur)
                elif op == "l" and cur:
                    x, y = P(float(o[0]), float(o[1])); cur["pts"].append(_pt(x, y))
                elif op in ("c", "v", "y") and cur:
                    v = [float(a) for a in o]
                    u = cur["pts"][-1]
                    if op == "c": c1, c2, p = P(v[0], v[1]), P(v[2], v[3]), P(v[4], v[5])
                    elif op == "v": c1, c2, p = u[:2], P(v[0], v[1]), P(v[2], v[3])
                    else: c1, c2, p = P(v[0], v[1]), P(v[2], v[3]), P(v[2], v[3])
                    u[4], u[5] = c1; cur["pts"].append([p[0], p[1], c2[0], c2[1], p[0], p[1]])
                elif op == "h" and cur:
                    pts = cur["pts"]
                    if len(pts) > 1 and abs(pts[0][0] - pts[-1][0]) < 1e-4 and abs(pts[0][1] - pts[-1][1]) < 1e-4:
                        pts[0][2], pts[0][3] = pts[-1][2], pts[-1][3]; pts.pop()
                    cur["fechado"] = True
                elif op == "re":
                    x, y, w, h = [float(a) for a in o]
                    cur = {"fechado": True, "pts": [_pt(*P(x, y)), _pt(*P(x + w, y)), _pt(*P(x + w, y + h)), _pt(*P(x, y + h))]}; subs.append(cur); cur = None
                elif op in ("S", "s", "f", "F", "f*", "B", "B*", "b", "b*", "n"):
                    if op in ("s", "b", "b*") and cur: cur["fechado"] = True
                    if op in ("f", "F", "f*", "B", "B*", "b", "b*"):
                        for s_ in subs: s_["fechado"] = True
                    emitir(op in ("f", "F", "f*", "B", "B*", "b", "b*"), op in ("S", "s", "B", "B*", "b", "b*"), "evenodd" if "*" in op else "nonzero")
                elif op in ("W", "W*"):
                    clip_pend = "evenodd" if op == "W*" else "nonzero"
                    for s_ in subs: s_["fechado"] = True
                elif op == "w": gs["lw"] = float(o[0])
                elif op == "J": gs["cap"] = int(o[0])
                elif op == "j": gs["join"] = int(o[0])
                elif op == "M": gs["miter"] = float(o[0])
                elif op == "d": gs["dash"] = ([float(v) for v in o[0]], float(o[1]))
                elif op in ("g", "G", "rg", "RG", "k", "K"):
                    cs = {"g": ("gray", None), "rg": ("rgb", None), "k": ("cmyk", None)}[op.lower()]
                    c = self._cor(cs, list(o), res)
                    if op.islower(): gs["fill"], gs["fcs"] = c, cs
                    else: gs["stroke"], gs["scs"] = c, cs
                elif op in ("cs", "CS"):
                    cs = self._cs(o[0], res)
                    ini = {"gray": [0], "rgb": [0, 0, 0], "cmyk": [0, 0, 0, 1], "sep": [1], "devn": [1]}.get(cs[0], [0])
                    c = self._cor(cs, ini, res) if cs[0] != "pattern" else None
                    if op == "cs": gs["fcs"], gs["fill"] = cs, c
                    else: gs["scs"], gs["stroke"] = cs, c
                elif op in ("sc", "scn", "SC", "SCN"):
                    cs = gs["fcs"] if op.islower() else gs["scs"]
                    if o and isinstance(o[-1], pikepdf.Name):   # padrão (degradê)
                        pats = (res or {}).get("/Pattern", {})
                        pat = pats.get(o[-1]) if pats else None
                        c = None
                        if pat is not None and int(pat.get("/PatternType", 0)) == 2:
                            c = self._gradiente(pat["/Shading"], mmul([float(v) for v in pat.get("/Matrix", [1, 0, 0, 1, 0, 0])], self.F), res)
                        else:
                            self.rel["padrao"] += 1
                            c = {"k": "cmyk", "v": [0, 0, 0, 30]}
                    else:
                        c = self._cor(cs, list(o), res)
                    if op.islower(): gs["fill"] = c
                    else: gs["stroke"] = c
                elif op == "gs":
                    eg = ((res or {}).get("/ExtGState") or {}).get(o[0])
                    if eg is not None:
                        if "/ca" in eg: gs["ca"] = float(eg["/ca"])
                        if "/CA" in eg: gs["CA"] = float(eg["/CA"])
                        if "/BM" in eg:
                            b = eg["/BM"]; b = b[0] if isinstance(b, pikepdf.Array) else b
                            gs["bm"] = bm_map.get(str(b), "normal")
                        if "/OP" in eg: gs["OP"] = bool(eg["/OP"]); gs["op"] = bool(eg.get("/op", eg["/OP"]))
                        if "/op" in eg: gs["op"] = bool(eg["/op"])
                        if "/LW" in eg: gs["lw"] = float(eg["/LW"])
                        if "/SMask" in eg and str(eg["/SMask"]) != "/None": self.rel["mascara_suave"] += 1
                elif op == "sh":
                    sh = ((res or {}).get("/Shading") or {}).get(o[0])
                    if sh is not None:
                        g = self._gradiente(sh, gs["ctm"], res)
                        if g:
                            area = gs["clip"] or [{"fechado": True, "pts": [_pt(*P(-1e4, -1e4)), _pt(*P(1e4, -1e4)), _pt(*P(1e4, 1e4)), _pt(*P(-1e4, 1e4))]}]
                            o2 = {"id": self.ids(), "tipo": "caminho", "subs": [dict(s, pts=[list(p) for p in s["pts"]]) for s in area], "regra": "nonzero", "preench": g, "traco": None}
                            if gs["ca"] < 1: o2["op"] = gs["ca"]
                            gs["cont"].append(o2)
                elif op == "Do":
                    xo = ((res or {}).get("/XObject") or {}).get(o[0])
                    if xo is None: continue
                    st = str(xo.get("/Subtype"))
                    if st == "/Form":
                        sv = dict(gs)
                        gs["ctm"] = mmul([float(v) for v in xo.get("/Matrix", [1, 0, 0, 1, 0, 0])], gs["ctm"])
                        self._stream(xo, xo.get("/Resources", res), gs, prof + 1)
                        gs = sv
                    elif st == "/Image":
                        self._imagem(pikepdf.PdfImage(xo), xo, gs)
                elif op == "INLINE IMAGE":
                    self._imagem(o[0], None, gs)
                elif op == "BDC" and len(o) >= 2 and str(o[0]) == "/OC":
                    nome = None
                    props = o[1] if isinstance(o[1], pikepdf.Dictionary) else ((res or {}).get("/Properties") or {}).get(o[1])
                    if props is not None:
                        ocg = props
                        if "/OCGs" in props:
                            ocg = props["/OCGs"]; ocg = ocg[0] if isinstance(ocg, pikepdf.Array) else ocg
                        nome = str(ocg.get("/Name", "")) or None
                    self.camada_pilha.append(nome or self.camada_pilha[-1])
                    if gs["raiz"] and nome: gs["cont"] = self._camada(nome)
                elif op == "BMC" or op == "BDC":
                    self.camada_pilha.append(self.camada_pilha[-1])
                elif op == "EMC":
                    if len(self.camada_pilha) > 1: self.camada_pilha.pop()
                    if gs["raiz"]: gs["cont"] = self._camada(self.camada_pilha[-1])
                elif op == "BT":
                    tm = tlm = [1, 0, 0, 1, 0, 0]
                elif op == "Tm":
                    tm = tlm = [float(v) for v in o]
                elif op in ("Td", "TD"):
                    if op == "TD": tl = -float(o[1])
                    tlm = mmul([1, 0, 0, 1, float(o[0]), float(o[1])], tlm); tm = tlm
                elif op == "TL": tl = float(o[0])
                elif op == "T*":
                    tlm = mmul([1, 0, 0, 1, 0, -tl], tlm); tm = tlm
                elif op in ("Tj", "TJ", "'", '"'):
                    if op in ("'", '"'):
                        tlm = mmul([1, 0, 0, 1, 0, -tl], tlm); tm = tlm
                    x, y = mpt(mmul(tm, gs["ctm"]), 0, 0)
                    self.textos.append((x, y, gs["fill"], gs["op"]))
            except Exception as e:
                self.rel["erros"].append(f"{op}: {str(e)[:80]}")

    def _imagem(self, pimg, xo, gs):
        # sem a máscara aplicada pelo pikepdf (ele converteria CMYK → RGB sem perfil): a máscara é lida aqui
        try:
            im = pimg.as_pil_image(apply_mask=False)
        except TypeError:
            im = pimg.as_pil_image()
        except Exception as e:
            self.rel["erros"].append(f"imagem: {str(e)[:80]}"); return
        ma = None
        if xo is not None and ("/SMask" in xo or isinstance(xo.get("/Mask"), pikepdf_Stream())):
            try:
                import pikepdf
                if "/SMask" in xo:
                    ma = pikepdf.PdfImage(xo["/SMask"]).as_pil_image().convert("L")
                else:   # máscara de estêncil: 1 = escondido (a não ser que /Decode [1 0])
                    mk = xo["/Mask"]; ma = pikepdf.PdfImage(mk).as_pil_image().convert("L")
                    dec = [float(v) for v in mk.get("/Decode", [0, 1])]
                    from PIL import ImageOps
                    if dec[0] < dec[1]: ma = ImageOps.invert(ma)
                if ma.size != im.size: ma = ma.resize(im.size)
            except Exception:
                ma = None
        if im.mode == "1":
            im = im.convert("L")
        if ma is not None and im.mode != "CMYK":
            im = im.convert("RGBA"); im.putalpha(ma)
        iid = _salvar_imagem(im, self.doc, self.ids, mascara=ma if im.mode == "CMYK" else None)
        W, H = im.size
        m = mmul([1 / W, 0, 0, -1 / H, 0, 1], gs["ctm"])
        o = {"id": self.ids(), "tipo": "imagem", "img": iid, "m": [round(v, 6) for v in m]}
        if gs["ca"] < 1: o["op"] = gs["ca"]
        gs["cont"].append(o)


def pikepdf_Stream():
    import pikepdf
    return pikepdf.Stream


def _limpar_clips(itens):
    """Grupo de corte cujo único conteúdo é um caminho igual ao corte (degradê do Illustrator) → só o caminho."""
    out = []
    for o in itens:
        if o.get("tipo") == "grupo":
            o["itens"] = _limpar_clips(o["itens"])
            if o.get("clip") and len(o["itens"]) == 2 and o["itens"][1].get("tipo") == "caminho" and o["itens"][1]["subs"] == o["itens"][0]["subs"]:
                out.append(o["itens"][1]); continue
            if o.get("clip") and len(o["itens"]) == 1:
                continue   # corte sem conteúdo
        out.append(o)
    return out


_PDF14 = {"helvetica": "Arial", "arialmt": "Arial", "times": "Times New Roman", "timesnewromanpsmt": "Times New Roman", "courier": "Courier New"}


def _fonte_nome(nome, doc_fontes=True):
    """nome PostScript do PDF → (família, estilo, achou). doc_fontes: vale a fonte embutida registrada (achou="embutida")."""
    nome = re.sub(r"^[A-Z]{6}\+", "", nome or "")
    base_, _, est_ = nome.partition("-")
    if base_.replace(",", "").lower() in _PDF14:   # 14 fontes padrão do PDF → equivalentes do Windows
        nome = _PDF14[base_.replace(",", "").lower()] + ("-" + est_ if est_ else "")
    from Functions import vetor_kanivete as vk
    est = vk._estilos()
    alvo = nome.replace(" ", "").lower()
    for fam, lista in est.items():
        for e in lista:
            if (e.get("ps") or "").replace(" ", "").lower() == alvo:
                return fam, e["estilo"], True
    base, _, estilo = nome.partition("-")
    for fam in est:
        if fam.replace(" ", "").lower() == base.lower():
            ests = [e["estilo"] for e in est[fam]]
            est_ok = next((s for s in ests if s.replace(" ", "").lower() == (estilo or "regular").lower()), None)
            return fam, est_ok or "Regular", True
    if doc_fontes:
        e = vk._fonte_doc(nome, "")
        if e and (e.get("ps") or "").replace(" ", "").lower() == alvo:
            return e["fam"], e["estilo"], "embutida"
    return nome or "Arial", "Regular", False


def _nome_fonte(tt, ps):
    """(família, estilo) pela tabela name (16/17 preferidos); senão pelo nome PostScript "Familia-Estilo"."""
    try:
        nm = tt["name"]
        fam = nm.getDebugName(16) or nm.getDebugName(1); est = nm.getDebugName(17) or nm.getDebugName(2)
        if fam and not re.match(r"^[A-Z]{6}\+", fam):
            if not est or "-" in est or " " not in est and len(est) > 14:   # instância de fonte variável: "ArchivoRoman-SemiBold"
                est = ps.partition("-")[2] or "Regular"
                est = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", est)
            return fam.strip(), est.strip()
    except Exception:
        pass
    base, _, est = ps.partition("-")
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", base), re.sub(r"(?<=[a-z])(?=[A-Z])", " ", est or "Regular")


def _cff_para_otf(buf):
    """CFF puro (FontFile3/Type1C, ex.: MyriadPro do Illustrator) → OTF com cmap pelos nomes dos glifos."""
    from fontTools.cffLib import CFFFontSet
    from fontTools.fontBuilder import FontBuilder
    from fontTools.ttLib import newTable
    from fontTools.agl import toUnicode
    cff = CFFFontSet(); cff.decompile(io.BytesIO(buf), None)
    nome = cff.fontNames[0]; top = cff[nome]
    if hasattr(top, "ROS"):
        raise ValueError("CFF CID")
    ordem = list(top.charset); cs = top.CharStrings
    upm = round(1 / top.FontMatrix[0]) if top.FontMatrix[0] else 1000
    cmap, hm = {}, {}
    for g in ordem:
        u = toUnicode(g)
        if len(u) == 1 and ord(u) not in cmap:
            cmap[ord(u)] = g
        c = cs[g]; c.decompile()
        try:
            b = c.calcBounds(cs)
        except Exception:
            b = None
        hm[g] = (int(round(getattr(c, "width", None) or top.Private.defaultWidthX)), int(b[0]) if b else 0)
    fb = FontBuilder(upm, isTTF=False)
    fb.setupGlyphOrder(ordem); fb.setupCharacterMap(cmap); fb.setupHorizontalMetrics(hm)
    fb.setupHorizontalHeader(ascent=int(upm * 0.75), descent=-int(upm * 0.25))
    ps = re.sub(r"^[A-Z]{6}\+", "", nome)
    fam, est = _nome_fonte(None, ps)
    fb.setupNameTable({"familyName": fam, "styleName": est, "psName": ps})
    fb.setupOS2(sTypoAscender=int(upm * 0.75), sTypoDescender=-int(upm * 0.25), usWinAscent=upm, usWinDescent=int(upm * 0.3))
    fb.setupPost(); fb.setupMaxp() if hasattr(fb, "setupMaxp") else None
    t = newTable("CFF "); t.cff = cff; cff.fontNames[0] = ps; fb.font["CFF "] = t
    out = io.BytesIO(); fb.font.save(out)
    return out.getvalue(), fam, est


def _fontes_embutidas(fz, doc, rel):
    """Fontes embutidas no PDF/AI que NÃO estão instaladas → arquivos (subconjuntos do mesmo nome juntados) registrados
    como fontes do documento (doc.fontes). O texto fica editável e igual ao original; letras fora do subconjunto somem."""
    from fontTools.ttLib import TTFont
    from Functions import vetor_kanivete as vk
    grupos, vistos = {}, set()
    for pg in fz:
        try:
            lst = pg.get_fonts()
        except Exception:
            continue
        for (xref, ext, tipo, base, *_r) in lst:
            if xref in vistos: continue
            vistos.add(xref)
            ps = re.sub(r"^[A-Z]{6}\+", "", base or "")
            if not ps or _fonte_nome(ps, doc_fontes=False)[2]:
                continue   # instalada
            try:
                _n, e, _t, buf = fz.extract_font(xref)
            except Exception:
                continue
            if buf and e in ("ttf", "otf", "cff"):
                grupos.setdefault(ps, []).append((e, buf))
    os.makedirs(PASTA, exist_ok=True)
    regs = []
    for ps, bufs in grupos.items():
        try:
            tts = [(e, b) for e, b in bufs if e in ("ttf", "otf")]
            if tts:
                fonts = [TTFont(io.BytesIO(b)) for _e, b in tts]
                base_ = max(fonts, key=lambda t: len(t.getGlyphOrder()))
                if "glyf" in base_:   # junta os subconjuntos (o Illustrator mantém os ids dos glifos)
                    gb = base_["glyf"]
                    for t in fonts:
                        if t is base_ or "glyf" not in t or t.getGlyphOrder() != base_.getGlyphOrder(): continue
                        for g in t.getGlyphOrder():
                            if gb[g].numberOfContours == 0 and t["glyf"][g].numberOfContours != 0:
                                gb[g] = t["glyf"][g]
                if not base_.getBestCmap():
                    rel["fonte_sem_cmap"] += 1; continue
                fam, est = _nome_fonte(base_, ps)
                out = io.BytesIO(); base_.save(out); dados, ext = out.getvalue(), ".ttf" if "glyf" in base_ else ".otf"
            else:
                dados, fam, est = _cff_para_otf(bufs[0][1]); ext = ".otf"
            arq = os.path.join(PASTA, f"fonte_{re.sub(r'[^A-Za-z0-9_-]', '_', ps)}_{hashlib.md5(dados).hexdigest()[:8]}{ext}")
            with open(arq, "wb") as f:
                f.write(dados)
            regs.append({"fam": fam, "estilo": est, "ps": ps, "arquivo": arq})
        except Exception as e:
            rel["erros"].append(f"fonte {ps}: {str(e)[:60]}")
    if regs:
        vk.registrar_fontes(regs)
        doc["fontes"] = regs
    return regs


def _ai_privado(pdf, limite=16 << 20):
    """Dados nativos do Illustrator (PieceInfo/Illustrator/Private, AIPrivateDataN): o começo do arquivo .ai em texto
    (cabeçalho, cores, pranchetas, camadas). Zstandard (AI 2020+) ou zlib (antigos). → bytes | None"""
    try:
        pv = pdf.pages[0].obj["/PieceInfo"]["/Illustrator"]["/Private"]
        n = int(pv["/NumBlock"])
        raw = b"".join(pv[f"/AIPrivateData{i}"].read_bytes() for i in range(1, n + 1))
    except Exception:
        return None
    try:
        if raw.startswith(b"%AI24_ZStandard_Data"):
            try:
                import zstandard
                with zstandard.ZstdDecompressor().stream_reader(io.BytesIO(raw[20:])) as r:   # só o começo (rápido)
                    return r.read(limite)
            except ImportError:
                from compression import zstd   # Python 3.14+
                return zstd.ZstdDecompressor().decompress(raw[20:], max_length=limite)
        if raw.startswith(b"%AI12_CompressedData"):
            return zlib.decompressobj().decompress(raw[20:], limite)
        return raw[:limite]
    except Exception:
        return None


def _ai_nativo(pdf, doc, rel, ids):
    """Do .ai nativo: nomes das pranchetas e das cores especiais (viram amostras, como no painel do Illustrator)."""
    d = _ai_privado(pdf)
    if not d:
        return
    def txt(b):
        b = re.sub(rb"\\([()\\])", rb"\1", b)
        try:
            return b.decode("utf-8")
        except UnicodeDecodeError:
            return b.decode("latin-1")
    nomes = [txt(n) for n in re.findall(rb"\(((?:[^()\\]|\\.)*)\) /UnicodeString \(Name\)", d)]
    ps = doc["pranchetas"]
    if nomes and len(nomes) >= len(ps):
        for p, n in zip(ps, nomes[:len(ps)]):
            if n.strip(): p["nome"] = n.strip()
        rel["ai_pranchetas"] += 1
    ja = {a.get("nome") for a in doc.setdefault("amostras", [])}
    for c, m, y, k, nome in re.findall(rb"%%(?:CMYKCustomColor:|\+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+) \(((?:[^()\\]|\\.)*)\)", d.split(b"%%EndComments", 1)[0]):
        nome = txt(nome)
        if nome in ja or nome in ("[Registration]", "[Registro]"):
            continue
        ja.add(nome)
        doc["amostras"].append({"id": ids("a"), "nome": nome,
                                "cor": {"k": "spot", "nome": nome, "v": [round(float(x) * 100, 2) for x in (c, m, y, k)], "tint": 100}})


def importar_pdf(caminho):
    import pikepdf, fitz
    nome = os.path.splitext(os.path.basename(caminho))[0]
    doc, ids = _doc(nome), _Ids()
    rel = Counter(); rel["erros"] = []
    try:
        pdf = pikepdf.open(caminho)
    except Exception as e:
        return {"success": False, "error": f"não abriu como PDF: {e}" + (" — salve o .ai com 'Criar arquivo compatível com PDF'" if caminho.lower().endswith(".ai") else "")}
    fz = fitz.open(caminho, filetype="pdf")
    leitor = _PDF(pdf, doc, ids, rel)
    embutidas = _fontes_embutidas(fz, doc, rel)
    x_ab = 0
    faltando = set()
    for i, page in enumerate(pdf.pages):
        caixa = [float(v) for v in (page.obj.get("/TrimBox") or page.obj.get("/CropBox") or page.obj.get("/MediaBox"))]
        x0, y0, x1, y1 = min(caixa[0], caixa[2]), min(caixa[1], caixa[3]), max(caixa[0], caixa[2]), max(caixa[1], caixa[3])
        w, h = x1 - x0, y1 - y0
        doc["pranchetas"].append({"id": ids("p"), "nome": f"Prancheta {i + 1}", "x": x_ab, "y": 0, "w": round(w, 3), "h": round(h, 3)})
        F = [1, 0, 0, -1, x_ab - x0, y1]
        rot = int(page.obj.get("/Rotate", 0)) % 360
        if rot:
            rel["pagina_girada"] += 1
        leitor.F = F; leitor.textos = []
        leitor.pagina(page, F)
        # textos (PyMuPDF): posição/fonte/tamanho; cor = do operador de texto mais perto (content stream)
        fp = fz[i]
        inv = ~fp.transformation_matrix   # espaço do fitz → espaço do PDF
        txt_cam = leitor._camada(leitor.camada_pilha[0] if leitor.camada_pilha else "Camada 1")
        d = fp.get_text("rawdict" if False else "dict")
        for b in d.get("blocks", []):
            for ln in b.get("lines", []):
                for sp in ln.get("spans", []):
                    t = sp.get("text", "")
                    if not t.strip():
                        continue
                    ox, oy = sp["origin"]
                    p = fitz.Point(ox, oy) * inv
                    X, Y = mpt(F, p.x, p.y)
                    dx, dy = ln.get("dir", (1, 0))
                    cor = None
                    if leitor.textos:
                        cand = min(leitor.textos, key=lambda q: (q[0] - X) ** 2 + (q[1] - Y) ** 2)
                        if (cand[0] - X) ** 2 + (cand[1] - Y) ** 2 < (sp["size"] * 3) ** 2: cor = cand[2]
                    if not cor:
                        c = sp.get("color", 0); cor = {"k": "rgb", "v": [(c >> 16) & 255, (c >> 8) & 255, c & 255]}
                        if cor["v"] == [0, 0, 0]: cor = {"k": "cmyk", "v": [0, 0, 0, 100]}
                    fam, estilo, ok = _fonte_nome(sp.get("font"))
                    if not ok: faltando.add(sp.get("font"))
                    ang = math.atan2(dy, dx)
                    cs, sn = math.cos(ang), math.sin(ang)
                    txt_cam.append({"id": ids(), "tipo": "texto", "conteudo": t, "fam": fam, "estilo": estilo, "tam": round(sp["size"], 3),
                                    "entrelinha": None, "track": 0, "alin": "esq", "caixa": None,
                                    "m": [round(cs, 6), round(sn, 6), round(-sn, 6), round(cs, 6), round(X, 3), round(Y, 3)],
                                    "preench": cor, "traco": None, **({"fonte_original": sp.get("font")} if not ok else {})})
        x_ab += w + GAP
    _ai_nativo(pdf, doc, rel, ids)
    for c in doc["camadas"]:
        c["itens"] = _limpar_clips(c["itens"])
    doc["camadas"] = [c for c in doc["camadas"] if c["itens"]] or [{"id": ids("c"), "nome": "Camada 1", "visivel": True, "trava": False, "itens": []}]
    if not leitor.cmyk:
        doc["modoCor"] = "rgb"
    texto_ai = caminho.lower().endswith(".ai") and any("saved without PDF Content" in fz[k].get_text() for k in range(min(1, len(fz))))
    relatorio = []
    if texto_ai: relatorio.append("Este .ai foi salvo SEM conteúdo PDF: no Illustrator, salve de novo marcando 'Criar arquivo compatível com PDF'.")
    if faltando: relatorio.append("Fontes não instaladas (trocadas por Arial; instale e reabra): " + ", ".join(sorted(f for f in faltando if f)))
    if embutidas: relatorio.append("Fontes não instaladas, usadas as EMBUTIDAS no arquivo (só as letras do original; para editar à vontade, instale): "
                                   + ", ".join(f"{e['fam']} {e['estilo']}" for e in embutidas))
    nomes = {"sombreamento": "degradês de malha/forma livre (só os lineares e radiais viram degradê)", "padrao": "padrões (pattern) viraram cinza 30%",
             "mascara_suave": "máscaras de opacidade (ignoradas)", "devn": "cores DeviceN convertidas", "pagina_girada": "páginas giradas (abrem sem a rotação)",
             "cmyk_com_alfa": "imagens CMYK com transparência (viraram RGB)", "fonte_sem_cmap": "fontes embutidas sem tabela de caracteres (trocadas por Arial)"}
    for k, txt in nomes.items():
        if rel[k]: relatorio.append(f"{rel[k]}× {txt}")
    if rel["erros"]: relatorio.append(f"{len(rel['erros'])} operações não lidas (ex.: {rel['erros'][0]})")
    return {"success": True, "doc": doc, "relatorio": relatorio}


# ═══════════════════════════ SVG ═══════════════════════════
_CORES_NOMES = {"black": "000000", "white": "ffffff", "red": "ff0000", "green": "008000", "blue": "0000ff", "yellow": "ffff00", "gray": "808080",
                "grey": "808080", "orange": "ffa500", "purple": "800080", "cyan": "00ffff", "magenta": "ff00ff", "lime": "00ff00", "navy": "000080",
                "silver": "c0c0c0", "maroon": "800000", "teal": "008080", "olive": "808000", "pink": "ffc0cb", "brown": "a52a2a"}
_HERDA = {"fill", "stroke", "stroke-width", "fill-rule", "font-family", "font-size", "font-weight", "font-style", "text-anchor", "stroke-linecap",
          "stroke-linejoin", "stroke-miterlimit", "stroke-dasharray", "stroke-dashoffset", "visibility", "fill-opacity", "stroke-opacity", "letter-spacing"}


def _num(v, padrao=0.0, ref=None):
    if v is None:
        return padrao
    m = re.match(r"\s*(-?[\d.]+(?:e-?\d+)?)\s*(px|pt|mm|cm|in|%)?", str(v))
    if not m:
        return padrao
    n = float(m.group(1)); u = m.group(2)
    if u == "%" and ref is not None: return n / 100 * ref
    return n * {"pt": 4 / 3, "mm": 96 / 25.4, "cm": 96 / 2.54, "in": 96}.get(u, 1)   # em px (usuário SVG)


def _svg_cor(v, op=1.0):
    if v is None: return None
    v = v.strip()
    if v in ("none", "transparent", ""): return None
    if v.startswith("url("): return ("url", re.search(r"#([^)'\"]+)", v).group(1))
    if v.startswith("#"):
        h = v[1:]
        if len(h) == 3: h = "".join(c * 2 for c in h)
        try: return {"k": "rgb", "v": [int(h[i:i + 2], 16) for i in (0, 2, 4)]}
        except ValueError: return None
    m = re.match(r"rgba?\(([^)]*)\)", v)
    if m:
        p = [x.strip() for x in m.group(1).split(",")]
        return {"k": "rgb", "v": [round(float(x[:-1]) * 2.55) if x.endswith("%") else int(float(x)) for x in p[:3]]}
    if m := re.match(r"device-cmyk\(([^)]*)\)", v):
        p = [float(x.strip().rstrip("%")) for x in re.split(r"[ ,]+", m.group(1).strip()) if x.strip()]
        return {"k": "cmyk", "v": [x * 100 if x <= 1 else x for x in p[:4]]}
    if v.lower() in _CORES_NOMES: return _svg_cor("#" + _CORES_NOMES[v.lower()])
    return {"k": "rgb", "v": [0, 0, 0]}


def _svg_matriz(t):
    m = [1, 0, 0, 1, 0, 0]
    for nome, args in re.findall(r"(\w+)\s*\(([^)]*)\)", t or ""):
        a = [float(x) for x in re.findall(r"-?[\d.]+(?:e-?\d+)?", args)]
        if nome == "matrix" and len(a) == 6: n = a
        elif nome == "translate": n = [1, 0, 0, 1, a[0], a[1] if len(a) > 1 else 0]
        elif nome == "scale": n = [a[0], 0, 0, a[1] if len(a) > 1 else a[0], 0, 0]
        elif nome == "rotate":
            r = math.radians(a[0]); c, s = math.cos(r), math.sin(r)
            n = [c, s, -s, c, 0, 0]
            if len(a) == 3: n = mmul(mmul([1, 0, 0, 1, -a[1], -a[2]], n), [1, 0, 0, 1, a[1], a[2]])
        elif nome == "skewX": n = [1, 0, math.tan(math.radians(a[0])), 1, 0, 0]
        elif nome == "skewY": n = [1, math.tan(math.radians(a[0])), 0, 1, 0, 0]
        else: continue
        m = mmul(n, m)   # SVG: a da esquerda é a mais externa
    return m


def _arco(x1, y1, rx, ry, phi, fa, fs, x2, y2):
    """Arco SVG → lista de cúbicas [(c1, c2, p)]."""
    if rx == 0 or ry == 0: return [((x1, y1), (x2, y2), (x2, y2))]
    rx, ry = abs(rx), abs(ry); ph = math.radians(phi); cp, sp = math.cos(ph), math.sin(ph)
    dx, dy = (x1 - x2) / 2, (y1 - y2) / 2
    x1p, y1p = cp * dx + sp * dy, -sp * dx + cp * dy
    lam = x1p ** 2 / rx ** 2 + y1p ** 2 / ry ** 2
    if lam > 1: rx, ry = rx * math.sqrt(lam), ry * math.sqrt(lam)
    num = rx ** 2 * ry ** 2 - rx ** 2 * y1p ** 2 - ry ** 2 * x1p ** 2
    co = math.sqrt(max(0, num / (rx ** 2 * y1p ** 2 + ry ** 2 * x1p ** 2))) * (-1 if fa == fs else 1)
    cxp, cyp = co * rx * y1p / ry, -co * ry * x1p / rx
    cx, cy = cp * cxp - sp * cyp + (x1 + x2) / 2, sp * cxp + cp * cyp + (y1 + y2) / 2
    ang = lambda ux, uy, vx, vy: math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)
    t1 = ang(1, 0, (x1p - cxp) / rx, (y1p - cyp) / ry)
    dt = ang((x1p - cxp) / rx, (y1p - cyp) / ry, (-x1p - cxp) / rx, (-y1p - cyp) / ry)
    if not fs and dt > 0: dt -= 2 * math.pi
    elif fs and dt < 0: dt += 2 * math.pi
    n = max(1, math.ceil(abs(dt) / (math.pi / 2))); d = dt / n; k = 4 / 3 * math.tan(d / 4)
    out = []
    E = lambda t: (cx + rx * math.cos(t) * cp - ry * math.sin(t) * sp, cy + rx * math.cos(t) * sp + ry * math.sin(t) * cp)
    dE = lambda t: (-rx * math.sin(t) * cp - ry * math.cos(t) * sp, -rx * math.sin(t) * sp + ry * math.cos(t) * cp)
    for i in range(n):
        a, b = t1 + i * d, t1 + (i + 1) * d
        p0, p3, d0, d3 = E(a), E(b), dE(a), dE(b)
        out.append(((p0[0] + k * d0[0], p0[1] + k * d0[1]), (p3[0] - k * d3[0], p3[1] - k * d3[1]), p3))
    return out


def svg_path(d):
    """Atributo d → subcaminhos (coordenadas do usuário SVG)."""
    toks = re.findall(r"[MmLlHhVvCcSsQqTtAaZz]|-?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?", d or "")
    subs, cur, i, cmd = [], None, 0, None
    x = y = sx = sy = 0; ultc = None; ultq = None

    def num():
        nonlocal i
        v = float(toks[i]); i += 1; return v

    def add(px, py):
        cur["pts"].append(_pt(px, py))

    def cub(c1, c2, p):
        u = cur["pts"][-1]; u[4], u[5] = c1; cur["pts"].append([p[0], p[1], c2[0], c2[1], p[0], p[1]])

    while i < len(toks):
        if re.match(r"[A-Za-z]", toks[i]):
            cmd = toks[i]; i += 1
            if cmd in "Zz":
                if cur:
                    pts = cur["pts"]
                    if len(pts) > 1 and abs(pts[0][0] - pts[-1][0]) < 1e-6 and abs(pts[0][1] - pts[-1][1]) < 1e-6:
                        pts[0][2], pts[0][3] = pts[-1][2], pts[-1][3]; pts.pop()
                    cur["fechado"] = True; x, y = sx, sy
                    cur = None
                continue
        if cmd is None: break
        rel = cmd.islower(); C = cmd.upper(); ox, oy = (x, y) if rel else (0, 0)
        if cur is None and C != "M":
            cur = {"fechado": False, "pts": [_pt(x, y)]}; subs.append(cur)
        if C == "M":
            x, y = num() + ox, num() + oy; sx, sy = x, y
            cur = {"fechado": False, "pts": [_pt(x, y)]}; subs.append(cur)
            cmd = "l" if rel else "L"; ultc = ultq = None; continue
        if C == "L": x, y = num() + ox, num() + oy; add(x, y); ultc = ultq = None
        elif C == "H": x = num() + ox; add(x, y); ultc = ultq = None
        elif C == "V": y = num() + oy; add(x, y); ultc = ultq = None
        elif C == "C":
            c1 = (num() + ox, num() + oy); c2 = (num() + ox, num() + oy); x, y = num() + ox, num() + oy
            cub(c1, c2, (x, y)); ultc = c2; ultq = None
        elif C == "S":
            c1 = (2 * x - ultc[0], 2 * y - ultc[1]) if ultc else (x, y)
            c2 = (num() + ox, num() + oy); x, y = num() + ox, num() + oy; cub(c1, c2, (x, y)); ultc = c2; ultq = None
        elif C in "QT":
            q = (num() + ox, num() + oy) if C == "Q" else ((2 * x - ultq[0], 2 * y - ultq[1]) if ultq else (x, y))
            px, py = num() + ox, num() + oy
            cub((x + 2 / 3 * (q[0] - x), y + 2 / 3 * (q[1] - y)), (px + 2 / 3 * (q[0] - px), py + 2 / 3 * (q[1] - py)), (px, py))
            x, y = px, py; ultq = q; ultc = None
        elif C == "A":
            rx, ry, phi = num(), num(), num(); fa, fs = int(num()), int(num()); px, py = num() + ox, num() + oy
            for c1, c2, p in _arco(x, y, rx, ry, phi, fa, fs, px, py): cub(c1, c2, p)
            x, y = px, py; ultc = ultq = None
        else:
            i += 1
    return [s for s in subs if len(s["pts"]) > 1 or s["fechado"]]


def importar_svg(caminho):
    from lxml import etree
    from PIL import Image
    import io
    nome = os.path.splitext(os.path.basename(caminho))[0]
    doc, ids = _doc(nome, "rgb"), _Ids()
    rel = Counter()
    raiz = etree.parse(caminho, etree.XMLParser(huge_tree=True, recover=True)).getroot()
    ns = lambda t: t.split("}")[-1] if isinstance(t, str) else ""
    por_id = {e.get("id"): e for e in raiz.iter() if e.get("id")}
    # CSS das <style>: .classe, #id, elemento (o que o Illustrator exporta)
    css = {}
    for st in raiz.iter("{*}style"):
        for sel, corpo in re.findall(r"([^{}]+)\{([^}]*)\}", re.sub(r"/\*.*?\*/", "", st.text or "", flags=re.S)):
            regras = dict((k.strip(), v.strip()) for k, _, v in (r.partition(":") for r in corpo.split(";")) if v.strip())
            for s in sel.split(","):
                css.setdefault(s.strip(), {}).update(regras)
    vb = [float(x) for x in re.findall(r"-?[\d.]+(?:e-?\d+)?", raiz.get("viewBox") or "")]
    W = _num(raiz.get("width"), None) if raiz.get("width") and "%" not in raiz.get("width") else None
    H = _num(raiz.get("height"), None) if raiz.get("height") and "%" not in raiz.get("height") else None
    if len(vb) == 4:
        sx = (W / vb[2]) if W else 1; sy = (H / vb[3]) if H else sx
        base = mmul([1, 0, 0, 1, -vb[0], -vb[1]], [sx * 0.75, 0, 0, sy * 0.75, 0, 0])   # px → pt
        aw, ah = (W or vb[2]) * 0.75, (H or vb[3] * sx) * 0.75
    else:
        base = [0.75, 0, 0, 0.75, 0, 0]; aw, ah = (W or 800) * 0.75, (H or 600) * 0.75
    doc["pranchetas"].append({"id": ids("p"), "nome": "Prancheta 1", "x": 0, "y": 0, "w": round(aw, 3), "h": round(ah, 3)})
    camada = {"id": ids("c"), "nome": "Camada 1", "visivel": True, "trava": False, "itens": []}
    doc["camadas"].append(camada)

    def estilo(e, herdado):
        s = {k: v for k, v in herdado.items() if k in _HERDA}
        for k in ("fill", "stroke", "stroke-width", "fill-rule", "font-family", "font-size", "font-weight", "font-style", "text-anchor", "opacity",
                  "fill-opacity", "stroke-opacity", "stroke-linecap", "stroke-linejoin", "stroke-miterlimit", "stroke-dasharray", "display",
                  "visibility", "clip-path", "mix-blend-mode", "letter-spacing", "stroke-dashoffset"):
            if e.get(k) is not None: s[k] = e.get(k)
        for sel in [ns(e.tag)] + ["." + c for c in (e.get("class") or "").split()] + (["#" + e.get("id")] if e.get("id") else []):
            s.update(css.get(sel, {}))
        for r in (e.get("style") or "").split(";"):
            k, _, v = r.partition(":")
            if v.strip(): s[k.strip()] = v.strip()
        return s

    def grad(gid, m, bbox):
        g = por_id.get(gid)
        if g is None: return None
        cad, at = g, {}
        while cad is not None:   # herança por href
            for k, v in cad.attrib.items():
                at.setdefault(k.split("}")[-1], v)
            stops = at.get("_stops") or [s for s in cad if ns(s.tag) == "stop"]
            if stops: at["_stops"] = stops
            href = cad.get("{http://www.w3.org/1999/xlink}href") or cad.get("href")
            cad = por_id.get(href[1:]) if href else None
        paradas = []
        for s in at.get("_stops", []):
            ss = estilo(s, {})
            c = _svg_cor(s.get("stop-color") or ss.get("stop-color") or "#000")
            off = s.get("offset", "0"); off = float(off[:-1]) / 100 if off.endswith("%") else float(off)
            paradas.append({"p": round(off, 4), "cor": c if isinstance(c, dict) else {"k": "rgb", "v": [0, 0, 0]}})
        if not paradas: return None
        gm = _svg_matriz(at.get("gradientTransform"))
        if at.get("gradientUnits") != "userSpaceOnUse":
            x0, y0, x1, y1 = bbox
            gm = mmul(gm, [x1 - x0 or 1, 0, 0, y1 - y0 or 1, x0, y0]); M = gm   # bbox já está no doc
            frac = lambda k, d: float(at.get(k, d).rstrip("%")) / (100 if str(at.get(k, d)).endswith("%") else 1)
        else:
            M = mmul(gm, m)
            frac = lambda k, d: _num(at.get(k, d))
        if ns(g.tag) == "linearGradient":
            a = mpt(M, frac("x1", "0"), frac("y1", "0")); b = mpt(M, frac("x2", "1" if at.get("gradientUnits") != "userSpaceOnUse" else "100%"), frac("y2", "0"))
            return {"k": "grad", "tipo": "lin", "a": list(a), "b": list(b), "paradas": paradas}
        c = mpt(M, frac("cx", "0.5"), frac("cy", "0.5")); f = mpt(M, frac("fx", at.get("cx", "0.5")), frac("fy", at.get("cy", "0.5")))
        return {"k": "grad", "tipo": "rad", "a": list(c), "f": list(f), "r": frac("r", "0.5") * mescala(M), "paradas": paradas}

    def pintar(o, s, m):
        bbox = _bbox_subs(o["subs"]) if o.get("subs") else (0, 0, 1, 1)
        def cor(k, opk):
            c = _svg_cor(s.get(k, "#000" if k == "fill" else None))
            if isinstance(c, tuple): c = grad(c[1], m, bbox)
            return c
        o["preench"] = cor("fill", "fill-opacity")
        sc = cor("stroke", "stroke-opacity")
        o["traco"] = {"cor": sc, "larg": round(_num(s.get("stroke-width"), 1) * mescala(m), 4), "cap": s.get("stroke-linecap", "butt"),
                      "junc": s.get("stroke-linejoin", "miter").replace("miter-clip", "miter").replace("arcs", "miter"),
                      "miter": float(s.get("stroke-miterlimit", 4)),
                      "tracejado": [round(_num(v) * mescala(m), 3) for v in re.split(r"[ ,]+", s.get("stroke-dasharray", "")) if v and v != "none"],
                      "fase": _num(s.get("stroke-dashoffset"))} if sc else None
        op = float(s.get("opacity", 1)) * float(s.get("fill-opacity", 1) if o["preench"] else 1)
        if op < 1: o["op"] = round(op, 3)
        if s.get("mix-blend-mode") == "multiply": o["bm"] = "multiplicacao"
        if s.get("fill-rule") == "evenodd": o["regra"] = "evenodd"
        return o

    def clipar(e_cp, m):
        subs = []
        for f in e_cp.iter():
            if ns(f.tag) in ("path", "rect", "circle", "ellipse", "polygon", "polyline"):
                subs += forma(f, mmul(_svg_matriz(f.get("transform")), m)) or []
        return {"id": ids(), "tipo": "caminho", "subs": subs, "regra": "nonzero", "preench": None, "traco": None}

    def forma(e, m):
        t = ns(e.tag)
        if t == "path": subs = svg_path(e.get("d"))
        elif t == "rect":
            w, h = _num(e.get("width")), _num(e.get("height"))
            rx = _num(e.get("rx"), _num(e.get("ry"), 0))
            if w <= 0 or h <= 0: return None
            subs = _ret_subs(_num(e.get("x")), _num(e.get("y")), w, h, rx)
        elif t == "circle":
            r = _num(e.get("r")); subs = _elipse_subs(_num(e.get("cx")), _num(e.get("cy")), r, r)
        elif t == "ellipse":
            subs = _elipse_subs(_num(e.get("cx")), _num(e.get("cy")), _num(e.get("rx")), _num(e.get("ry")))
        elif t == "line":
            subs = [{"fechado": False, "pts": [_pt(_num(e.get("x1")), _num(e.get("y1"))), _pt(_num(e.get("x2")), _num(e.get("y2")))]}]
        elif t in ("polygon", "polyline"):
            v = [float(x) for x in re.findall(r"-?[\d.]+(?:e-?\d+)?", e.get("points") or "")]
            subs = [{"fechado": t == "polygon", "pts": [_pt(v[i], v[i + 1]) for i in range(0, len(v) - 1, 2)]}]
        else:
            return None
        return _transformar_subs(subs, m)

    def andar(e, herdado, m, cont):
        t = ns(e.tag)
        if t in ("defs", "clipPath", "mask", "symbol", "marker", "pattern", "linearGradient", "radialGradient", "style", "title", "desc", "metadata", "filter"):
            if t in ("mask", "filter", "pattern") and e.getparent() is not None and ns(e.getparent().tag) != "defs": rel[t] += 1
            return
        s = estilo(e, herdado)
        if s.get("display") == "none" or s.get("visibility") == "hidden": return
        m = mmul(_svg_matriz(e.get("transform")), m)
        if e.get("mask") or s.get("mask"): rel["mask"] += 1
        if e.get("filter") or s.get("filter"): rel["filter"] += 1
        cp = s.pop("clip-path", None)
        if cp and "url(" in cp:
            alvo = por_id.get(re.search(r"#([^)'\"]+)", cp).group(1))
            if alvo is not None:
                g = {"id": ids(), "tipo": "grupo", "clip": True, "itens": [clipar(alvo, mmul(_svg_matriz(alvo.get("transform")), m))]}
                cont.append(g); cont = g["itens"]
        if t in ("svg", "g", "a", "switch") and e is not raiz:
            g = {"id": ids(), "tipo": "grupo", "itens": []}
            if e.get("id") and not re.match(r"^(Layer|Camada|_x)", e.get("id") or ""): g["nome"] = e.get("id")
            op = float(s.pop("opacity", 1))
            for f in e: andar(f, s, m, g["itens"])
            if op < 1:
                for f in g["itens"]: f["op"] = round(f.get("op", 1) * op, 3)
            if g["itens"]: cont.append(g)
            return
        if e is raiz:
            for f in e: andar(f, s, m, cont)
            return
        if t == "use":
            href = e.get("{http://www.w3.org/1999/xlink}href") or e.get("href")
            alvo = por_id.get((href or "#")[1:])
            if alvo is not None:
                mm = mmul([1, 0, 0, 1, _num(e.get("x")), _num(e.get("y"))], m)
                if ns(alvo.tag) == "symbol":
                    for f in alvo: andar(f, s, mm, cont)
                else: andar(alvo, s, mm, cont)
            return
        if t == "text":
            pedacos = []
            x0, y0 = _num(e.get("x")), _num(e.get("y"))
            def coleta(el, st, x, y):
                if el.text and el.text.strip(): pedacos.append((el.text, st, x, y))
                for f in el:
                    sf = estilo(f, st)
                    nx = _num(f.get("x"), None) if f.get("x") else None; ny = _num(f.get("y"), None) if f.get("y") else None
                    dxv, dyv = _num(f.get("dx")), _num(f.get("dy"))
                    coleta(f, sf, (nx if nx is not None else x) + dxv if nx is not None or dxv else None, (ny if ny is not None else y) + dyv if ny is not None or dyv else None)
                    if f.tail and f.tail.strip(): pedacos.append((f.tail, st, None, None))
            coleta(e, s, x0, y0)
            ux, uy = x0, y0
            for txt, st, x, y in pedacos:
                if x is not None: ux = x
                if y is not None: uy = y
                fam = (st.get("font-family") or "Arial").split(",")[0].strip().strip("'\"")
                fam_ok, estilo_f, ok = _fonte_nome(fam)
                peso = str(st.get("font-weight", "400"))
                if peso in ("bold", "700", "800", "900") and estilo_f == "Regular": estilo_f = "Bold"
                if not ok: rel["fonte:" + fam] += 1
                tam = _num(st.get("font-size"), 16)
                mm = mmul([1, 0, 0, 1, ux, uy], m)
                sc = mescala(mm)
                o = {"id": ids(), "tipo": "texto", "conteudo": txt.strip("\n"), "fam": fam_ok, "estilo": estilo_f, "tam": round(tam * sc, 3), "entrelinha": None,
                     "track": round(_num(st.get("letter-spacing")) / tam * 1000) if st.get("letter-spacing") else 0,
                     "alin": {"middle": "centro", "end": "dir"}.get(st.get("text-anchor"), "esq"), "caixa": None,
                     "m": [round(v / sc, 6) for v in mm[:4]] + [round(mm[4], 3), round(mm[5], 3)]}
                pintar(o, st, m); o.pop("regra", None)
                cont.append(o)
            return
        if t == "image":
            href = e.get("{http://www.w3.org/1999/xlink}href") or e.get("href") or ""
            try:
                if href.startswith("data:"):
                    im = Image.open(io.BytesIO(base64.b64decode(href.split(",", 1)[1])))
                else:
                    im = Image.open(os.path.join(os.path.dirname(caminho), href))
                im.load()
            except Exception:
                rel["imagem_faltando"] += 1; return
            iid = _salvar_imagem(im, doc, ids)
            w, h = _num(e.get("width"), im.width), _num(e.get("height"), im.height)
            mm = mmul([w / im.width, 0, 0, h / im.height, _num(e.get("x")), _num(e.get("y"))], m)
            cont.append({"id": ids(), "tipo": "imagem", "img": iid, "m": [round(v, 6) for v in mm]})
            return
        subs = forma(e, m)
        if subs:
            o = pintar({"id": ids(), "tipo": "caminho", "subs": subs, "regra": "nonzero"}, s, m)
            if o["preench"] or o["traco"]: cont.append(o)
        elif t not in ("path", "rect", "circle", "ellipse", "line", "polygon", "polyline"):
            rel["elemento:" + t] += 1

    andar(raiz, {"fill": "#000"}, base, camada["itens"])
    # grupos-camada do Illustrator (<g id="Layer_1">) no topo viram camadas
    if len(camada["itens"]) > 1 and all(o["tipo"] == "grupo" and o.get("nome") for o in camada["itens"]):
        doc["camadas"] = [{"id": ids("c"), "nome": o["nome"].replace("_", " "), "visivel": True, "trava": False, "itens": o["itens"]} for o in camada["itens"]]
    relatorio = []
    for k, n in rel.items():
        if k.startswith("fonte:"): relatorio.append(f"fonte não instalada: {k[6:]} (Arial no lugar)")
        elif k.startswith("elemento:"): relatorio.append(f"{n}× <{k[9:]}> não importado")
        else: relatorio.append(f"{n}× {dict(mask='máscara', filter='filtro (efeito)', pattern='padrão', imagem_faltando='imagem não encontrada').get(k, k)} ignorado")
    relatorio.append("SVG é RGB: para gráfica, use Fechamento › Converter cores para CMYK.")
    return {"success": True, "doc": doc, "relatorio": relatorio}


# ═══════════════════════════ PPTX ═══════════════════════════
EMU = 12700


def importar_pptx(caminho):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.enum.dml import MSO_FILL, MSO_THEME_COLOR
    from lxml import etree
    from PIL import Image
    import io
    nome = os.path.splitext(os.path.basename(caminho))[0]
    doc, ids = _doc(nome, "rgb"), _Ids()
    rel = Counter()
    prs = Presentation(caminho)
    SW, SH = prs.slide_width / EMU, prs.slide_height / EMU
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    # tema: cores e fontes
    tema_cores, tema_fontes = {}, {"major": "Calibri Light", "minor": "Calibri"}
    try:
        mestre = prs.slide_masters[0]
        for r in mestre.part.rels.values():
            if r.reltype.endswith("/theme"):
                x = etree.fromstring(r.target_part.blob)
                for el in x.iter(A + "clrScheme"):
                    for c in el:
                        v = c[0].get("lastClr") or c[0].get("val")   # sysClr: val="window", a cor está em lastClr
                        tema_cores[c.tag.replace(A, "")] = v
                for k in ("major", "minor"):
                    f = next(x.iter(A + k + "Font"), None)
                    if f is not None and f.find(A + "latin") is not None: tema_fontes[k] = f.find(A + "latin").get("typeface")
    except Exception:
        pass
    MAPA_T = {MSO_THEME_COLOR.TEXT_1: "dk1", MSO_THEME_COLOR.BACKGROUND_1: "lt1", MSO_THEME_COLOR.TEXT_2: "dk2", MSO_THEME_COLOR.BACKGROUND_2: "lt2",
              MSO_THEME_COLOR.DARK_1: "dk1", MSO_THEME_COLOR.LIGHT_1: "lt1", MSO_THEME_COLOR.DARK_2: "dk2", MSO_THEME_COLOR.LIGHT_2: "lt2",
              **{getattr(MSO_THEME_COLOR, f"ACCENT_{i}"): f"accent{i}" for i in range(1, 7)}}

    def cor(cf):
        try:
            if cf.type is None: return None
            if cf.type == 1:   # RGB
                h = str(cf.rgb)
            else:
                h = tema_cores.get(MAPA_T.get(cf.theme_color), "000000")
            v = [int(h[i:i + 2], 16) for i in (0, 2, 4)]
            b = cf.brightness or 0
            if b > 0: v = [round(x + (255 - x) * b) for x in v]
            elif b < 0: v = [round(x * (1 + b)) for x in v]
            return {"k": "rgb", "v": v}
        except Exception:
            return None

    def preench(sh):
        try:
            f = sh.fill
            if f.type == MSO_FILL.SOLID: return cor(f.fore_color)
            if f.type == MSO_FILL.GRADIENT:
                ps = [{"p": round(s.position, 4), "cor": cor(s.color) or {"k": "rgb", "v": [0, 0, 0]}} for s in f.gradient_stops]
                return {"k": "grad", "tipo": "lin", "_ang": f.gradient_angle or 0, "paradas": ps}
            if f.type is None:   # sem preenchimento próprio: herda do estilo da forma (<p:style><a:fillRef><a:schemeClr val="accent1">)
                return estilo_ref(sh, "fillRef")
        except Exception:
            pass
        return None

    def estilo_ref(sh, ref):
        el = sh._element.find(".//" + A + ref)
        if el is None or el.get("idx") == "0": return None
        sc = el.find(A + "schemeClr"); rgb = el.find(A + "srgbClr")
        h = rgb.get("val") if rgb is not None else tema_cores.get(sc.get("val") if sc is not None else "accent1", "4472C4")
        v = [int(h[i:i + 2], 16) for i in (0, 2, 4)]
        if sc is not None and sc.find(A + "shade") is not None: v = [round(x * int(sc.find(A + "shade").get("val")) / 100000) for x in v]
        return {"k": "rgb", "v": v}

    def linha(sh):
        try:
            ln = sh.line
            if ln.fill.type is None and not ln.width and sh._element.find(".//" + A + "lnRef") is not None and sh._element.find(".//" + A + "prstGeom") is not None:
                c = estilo_ref(sh, "lnRef")
                return {"cor": c, "larg": 1.0, "cap": "butt", "junc": "miter", "miter": 4, "tracejado": [], "fase": 0} if c else None
            if ln.fill.type in (None, MSO_FILL.BACKGROUND) and not ln.width: return None
            if ln.fill.type == MSO_FILL.BACKGROUND: return None
            c = cor(ln.color) if ln.fill.type == MSO_FILL.SOLID else None
            if not c: return None
            return {"cor": c, "larg": round((ln.width or 9525) / EMU, 3), "cap": "butt", "junc": "miter", "miter": 4, "tracejado": [], "fase": 0}
        except Exception:
            return None

    def geometria(sh, x, y, w, h):
        el = sh._element
        prst = el.find(".//" + A + "prstGeom")
        cust = el.find(".//" + A + "custGeom")
        if cust is not None:
            subs = []
            for p in cust.iter(A + "path"):
                pw, ph = float(p.get("w") or w * EMU) or 1, float(p.get("h") or h * EMU) or 1
                sx_, sy_ = w / (pw / EMU) / EMU if pw else 1, h / (ph / EMU) / EMU if ph else 1
                cur = None
                for c in p:
                    tag = c.tag.replace(A, "")
                    pts = [(x + float(q.get("x")) * sx_, y + float(q.get("y")) * sy_) for q in c.iter(A + "pt")]
                    if tag == "moveTo": cur = {"fechado": False, "pts": [_pt(*pts[0])]}; subs.append(cur)
                    elif tag == "lnTo" and cur: cur["pts"].append(_pt(*pts[0]))
                    elif tag == "cubicBezTo" and cur:
                        u = cur["pts"][-1]; u[4], u[5] = pts[0]; cur["pts"].append([pts[2][0], pts[2][1], pts[1][0], pts[1][1], pts[2][0], pts[2][1]])
                    elif tag == "quadBezTo" and cur:
                        u = cur["pts"][-1]; q, e = pts[0], pts[1]
                        u[4], u[5] = u[0] + 2 / 3 * (q[0] - u[0]), u[1] + 2 / 3 * (q[1] - u[1])
                        cur["pts"].append([e[0], e[1], e[0] + 2 / 3 * (q[0] - e[0]), e[1] + 2 / 3 * (q[1] - e[1]), e[0], e[1]])
                    elif tag == "close" and cur: cur["fechado"] = True
                    elif tag == "arcTo": rel["arco_livre"] += 1
            return subs
        g = prst.get("prst") if prst is not None else "rect"
        if g == "rect" or g is None: return _ret_subs(x, y, w, h)
        if g in ("roundRect", "snipRoundRect"):
            adj = 16667
            for gd in el.iter(A + "gd"):
                m_ = re.search(r"val (\d+)", gd.get("fmla") or "")
                if m_: adj = int(m_.group(1))
            return _ret_subs(x, y, w, h, min(w, h) * adj / 100000)
        if g == "ellipse": return _elipse_subs(x + w / 2, y + h / 2, w / 2, h / 2)
        if g == "triangle": return [{"fechado": True, "pts": [_pt(x + w / 2, y), _pt(x + w, y + h), _pt(x, y + h)]}]
        if g == "rtTriangle": return [{"fechado": True, "pts": [_pt(x, y), _pt(x + w, y + h), _pt(x, y + h)]}]
        if g == "diamond": return [{"fechado": True, "pts": [_pt(x + w / 2, y), _pt(x + w, y + h / 2), _pt(x + w / 2, y + h), _pt(x, y + h / 2)]}]
        if g in ("line", "straightConnector1"): return [{"fechado": False, "pts": [_pt(x, y), _pt(x + w, y + h)]}]
        if g == "parallelogram": d = w * 0.25; return [{"fechado": True, "pts": [_pt(x + d, y), _pt(x + w, y), _pt(x + w - d, y + h), _pt(x, y + h)]}]
        if g in ("hexagon", "pentagon", "octagon"):
            n = {"hexagon": 6, "pentagon": 5, "octagon": 8}[g]
            return [{"fechado": True, "pts": [_pt(x + w / 2 + w / 2 * math.cos(-math.pi / 2 + 2 * math.pi * i / n), y + h / 2 + h / 2 * math.sin(-math.pi / 2 + 2 * math.pi * i / n)) for i in range(n)]}]
        rel["forma:" + g] += 1
        return _ret_subs(x, y, w, h)

    def girar(m_ou_subs, sh, cx, cy, eh_subs=True):
        rot = getattr(sh, "rotation", 0) or 0
        xfrm = sh._element.find(".//" + A + "xfrm")
        fh = xfrm is not None and xfrm.get("flipH") == "1"; fv = xfrm is not None and xfrm.get("flipV") == "1"
        if not rot and not fh and not fv: return m_ou_subs
        r = math.radians(rot)
        M = mmul(mmul([1, 0, 0, 1, -cx, -cy], [(-1 if fh else 1), 0, 0, (-1 if fv else 1), 0, 0]), mmul([math.cos(r), math.sin(r), -math.sin(r), math.cos(r), 0, 0], [1, 0, 0, 1, cx, cy]))
        return _transformar_subs(m_ou_subs, M) if eh_subs else mmul(m_ou_subs, M)

    def texto(sh, x, y, w, h, cont):
        tf = sh.text_frame
        if not tf.text.strip(): return
        bp = tf._txBody.find(A + "bodyPr")
        ins = lambda k, d: (int(bp.get(k)) / EMU if bp is not None and bp.get(k) else d)
        l_, t_, r_ = ins("lIns", 7.2), ins("tIns", 3.6), ins("rIns", 7.2)
        run = next((r for p in tf.paragraphs for r in p.runs if r.text.strip()), None)
        f = run.font if run is not None else None
        fam = (f.name if f is not None and f.name else None) or tema_fontes["minor"]
        if fam.startswith("+mj"): fam = tema_fontes["major"]
        elif fam.startswith("+mn"): fam = tema_fontes["minor"]
        fam_ok, est, ok = _fonte_nome(fam)
        if not ok: rel["fonte:" + fam] += 1
        if f is not None and f.bold: est = "Bold" if est == "Regular" else est
        tam = f.size.pt if f is not None and f.size else (44 if sh.is_placeholder and "itle" in (sh.name or "") else 18)
        c = (cor(f.color) if f is not None and f.color and f.color.type is not None else None) or estilo_ref(sh, "fontRef") or {"k": "rgb", "v": [0, 0, 0]}
        al = tf.paragraphs[0].alignment
        alin = {2: "centro", 3: "dir", 4: "just"}.get(int(al) if al is not None else 1, "esq")
        cont.append({"id": ids(), "tipo": "texto", "conteudo": "\n".join(p.text for p in tf.paragraphs), "fam": fam_ok, "estilo": est, "tam": round(tam, 2),
                     "entrelinha": None, "track": 0, "alin": alin, "caixa": round(max(10, w - l_ - r_), 2),
                     "m": girar([1, 0, 0, 1, x + l_, y + t_], sh, x + w / 2, y + h / 2, False), "preench": c, "traco": None})

    def andar(shapes, M, cont):
        for sh in shapes:
            try:
                if sh.left is None: continue
                x, y = mpt(M, sh.left, sh.top)
                w, h = sh.width * M[0], sh.height * M[3]
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    xf = sh._element.grpSpPr.find(A + "xfrm")
                    ch_o, ch_e = xf.find(A + "chOff"), xf.find(A + "chExt")
                    cx, cy, cw, chh = int(ch_o.get("x")), int(ch_o.get("y")), int(ch_e.get("cx")) or 1, int(ch_e.get("cy")) or 1
                    sx_, sy_ = sh.width / cw, sh.height / chh
                    Mg = mmul([sx_, 0, 0, sy_, sh.left - cx * sx_, sh.top - cy * sy_], M)
                    g = {"id": ids(), "tipo": "grupo", "nome": sh.name, "itens": []}
                    andar(sh.shapes, Mg, g["itens"]); cont.append(g); continue
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(sh, "image", None) is not None and sh.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER:
                    im = Image.open(io.BytesIO(sh.image.blob)); im.load()
                    if any(getattr(sh, k, 0) for k in ("crop_left", "crop_right", "crop_top", "crop_bottom")): rel["corte_imagem"] += 1
                    iid = _salvar_imagem(im, doc, ids)
                    m = girar([w / im.width, 0, 0, h / im.height, x, y], sh, x + w / 2, y + h / 2, False)
                    cont.append({"id": ids(), "tipo": "imagem", "img": iid, "m": [round(v, 6) for v in m]}); continue
                if sh.has_chart if hasattr(sh, "has_chart") else False: rel["grafico"] += 1; continue
                if getattr(sh, "has_table", False) and sh.has_table: rel["tabela"] += 1; continue
                if sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER) or True:
                    p, t = preench(sh) if hasattr(sh, "fill") else None, linha(sh) if hasattr(sh, "line") else None
                    if p or t:
                        subs = girar(geometria(sh, x, y, w, h), sh, x + w / 2, y + h / 2)
                        if p and p.get("k") == "grad":
                            ang = math.radians(p.pop("_ang", 0)); cx, cy = x + w / 2, y + h / 2; L = (abs(w * math.cos(ang)) + abs(h * math.sin(ang))) / 2
                            p["a"] = [cx - L * math.cos(ang), cy + L * math.sin(ang)]; p["b"] = [cx + L * math.cos(ang), cy - L * math.sin(ang)]
                        cont.append({"id": ids(), "tipo": "caminho", "nome": sh.name, "subs": subs, "regra": "nonzero", "preench": p, "traco": t})
                    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                        texto(sh, x, y, w, h, cont)
            except Exception as e:
                rel["erro:" + str(e)[:60]] += 1

    for i, sl in enumerate(prs.slides):
        ox = i * (SW + GAP)
        doc["pranchetas"].append({"id": ids("p"), "nome": f"Slide {i + 1}", "x": ox, "y": 0, "w": round(SW, 3), "h": round(SH, 3)})
        M = [1 / EMU, 0, 0, 1 / EMU, ox, 0]
        cam = {"id": ids("c"), "nome": f"Slide {i + 1}", "visivel": True, "trava": False, "itens": []}
        try:
            bg = sl.background.fill
            if bg.type == MSO_FILL.SOLID:
                cam["itens"].append({"id": ids(), "tipo": "caminho", "nome": "Fundo", "subs": _ret_subs(ox, 0, SW, SH), "regra": "nonzero", "preench": cor(bg.fore_color), "traco": None})
        except Exception:
            pass
        andar(sl.shapes, M, cam["itens"])
        doc["camadas"].append(cam)
    relatorio = []
    for k, n in rel.items():
        if k.startswith("fonte:"): relatorio.append(f"fonte não instalada: {k[6:]} (Arial no lugar)")
        elif k.startswith("forma:"): relatorio.append(f"{n}× forma '{k[6:]}' virou retângulo")
        elif k.startswith("erro:"): relatorio.append(f"{n}× erro: {k[5:]}")
        else: relatorio.append(f"{n}× " + {"grafico": "gráfico não importado", "tabela": "tabela não importada", "corte_imagem": "imagem cortada (veio inteira)",
                                         "arco_livre": "arco em forma livre ignorado"}.get(k, k))
    relatorio.append("PowerPoint é RGB: para gráfica, use Fechamento › Converter cores para CMYK.")
    return {"success": True, "doc": doc, "relatorio": relatorio}


def ghostscript():
    """gswin64c/gswin32c/gs: PATH, Program Files/gs/*/bin ou %APPDATA%/CaniveteDoPailer/ghostscript (licença AGPL: não vai no app)."""
    import glob, shutil
    for n in ("gswin64c", "gswin32c", "gs"):
        p = shutil.which(n)
        if p: return p
    cands = []
    for raiz in (os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
                 os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), "CaniveteDoPailer", "ghostscript")):
        cands += glob.glob(os.path.join(raiz, "gs", "*", "bin", "gswin*c.exe")) + glob.glob(os.path.join(raiz, "bin", "gswin*c.exe"))
    return sorted(cands)[-1] if cands else None


GS_RELEASE = "gs10080"   # fallback se a API do GitHub não responder


def baixar_ghostscript():
    """Baixa o Ghostscript oficial (Artifex, AGPL — por isso não vai no app) para %APPDATA%/CaniveteDoPailer/ghostscript.
    Com 7-Zip na máquina só extrai o instalador (sem admin); sem 7-Zip roda o instalador silencioso (/S /D=pasta)."""
    import glob, json as _json, shutil, subprocess, urllib.request
    gs = ghostscript()
    if gs: return {"success": True, "gs": gs, "ja_tinha": True}
    dest = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), "CaniveteDoPailer", "ghostscript")
    os.makedirs(dest, exist_ok=True)
    url = f"https://github.com/ArtifexSoftware/ghostpdl-downloads/releases/download/{GS_RELEASE}/{GS_RELEASE}w64.exe"
    try:
        with urllib.request.urlopen("https://api.github.com/repos/ArtifexSoftware/ghostpdl-downloads/releases/latest", timeout=15) as r:
            url = next(a_["browser_download_url"] for a_ in _json.load(r)["assets"] if a_["name"].endswith("w64.exe"))
    except Exception:
        pass
    exe = os.path.join(dest, "_instalador_gs.exe")
    try:
        with urllib.request.urlopen(url, timeout=60) as r, open(exe, "wb") as f:
            shutil.copyfileobj(r, f, 1 << 20)
        sete = shutil.which("7z") or next((p for p in (os.path.join(os.environ.get(v, ""), "7-Zip", "7z.exe") for v in ("ProgramFiles", "ProgramW6432", "ProgramFiles(x86)")) if os.path.isfile(p)), None)
        sem_janela = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        if sete:
            subprocess.run([sete, "x", "-y", "-o" + dest, exe, "bin", "lib", "Resource", "iccprofiles"], capture_output=True, timeout=300, creationflags=sem_janela)
        if not glob.glob(os.path.join(dest, "bin", "gswin*c.exe")):
            subprocess.run(f'"{exe}" /S /D={dest}', timeout=600)
    except Exception as e:
        return {"success": False, "error": f"não baixou o Ghostscript: {e}"}
    finally:
        try: os.remove(exe)
        except OSError: pass
    gs = ghostscript()
    return {"success": True, "gs": gs} if gs else {"success": False, "error": "o Ghostscript baixou mas não instalou; instale em ghostscript.com/releases"}


def importar_eps(caminho):
    """EPS/PS → PDF pelo Ghostscript (EPSCrop: a caixa do desenho) → importador de PDF (cores, especiais, texto)."""
    import subprocess, tempfile
    gs = ghostscript()
    if not gs:
        return {"success": False, "sem_ghostscript": True, "error": "Para abrir EPS o Vetor usa o Ghostscript (gratuito), que não está instalado. Instale em "
                "ghostscript.com/releases (Windows 64 bits) e abra de novo — ou peça ao cliente o arquivo em PDF ou .ai compatível com PDF."}
    pdf = os.path.join(tempfile.gettempdir(), "vetor_kanivete", os.path.splitext(os.path.basename(caminho))[0] + "_eps.pdf")
    os.makedirs(os.path.dirname(pdf), exist_ok=True)
    r = subprocess.run([gs, "-dNOPAUSE", "-dBATCH", "-dSAFER", "-dQUIET", "-sDEVICE=pdfwrite", "-dEPSCrop", "-dPDFSETTINGS=/prepress",
                        "-dAutoRotatePages=/None", "-sColorConversionStrategy=LeaveColorUnchanged", f"-sOutputFile={pdf}", caminho],
                       capture_output=True, text=True, timeout=180, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if r.returncode != 0 or not os.path.isfile(pdf):
        return {"success": False, "error": "o Ghostscript não converteu o EPS: " + (r.stderr or r.stdout or "")[-300:]}
    res = importar_pdf(pdf)
    if res.get("success"):
        res["relatorio"] = ["EPS convertido pelo Ghostscript (PostScript → PDF): confira cores especiais e textos"] + (res.get("relatorio") or [])
        if res.get("doc"): res["doc"]["nome"] = os.path.splitext(os.path.basename(caminho))[0]
    return res


def importar(caminho):
    ext = os.path.splitext(caminho)[1].lower()
    if ext in (".eps", ".ps", ".epsf"):
        return importar_eps(caminho)
    if ext in (".pdf", ".ai"):
        return importar_pdf(caminho)
    if ext in (".svg", ".svgz"):
        return importar_svg(caminho)
    if ext == ".pptx":
        return importar_pptx(caminho)
    if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".psd"):
        from PIL import Image
        im = Image.open(caminho)
        doc, ids = _doc(os.path.splitext(os.path.basename(caminho))[0]), _Ids()
        dpi = (im.info.get("dpi") or (300, 300))[0] or 300
        w, h = im.width * 72 / dpi, im.height * 72 / dpi
        doc["pranchetas"].append({"id": ids("p"), "nome": "Prancheta 1", "x": 0, "y": 0, "w": w, "h": h})
        doc["imagens"]["i1"] = {"arquivo": os.path.abspath(caminho), "w": im.width, "h": im.height, "modo": im.mode, "nome": os.path.basename(caminho)}
        doc["camadas"].append({"id": "c1", "nome": "Camada 1", "visivel": True, "trava": False,
                               "itens": [{"id": ids(), "tipo": "imagem", "img": "i1", "m": [72 / dpi, 0, 0, 72 / dpi, 0, 0]}]})
        return {"success": True, "doc": doc, "relatorio": []}
    return {"success": False, "error": f"formato não suportado: {ext} (abre .aknv, .pdf, .ai, .svg, .pptx e imagens)"}
