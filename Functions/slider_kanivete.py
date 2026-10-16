"""slider_kanivete.py — Slider Kanivete (Sd, verde-água): apresentações 16:9 (PowerPoint/Canva/Figma no Kanivete).

Projeto .dknv = zip:
  documento.json — {formato: 'dknv', versao, nome, largura, altura, slides: [{id, fundo, transicao, dur, notas, oculto,
                    camadas: [{id, tipo: texto|forma|imagem|video, x, y, w, h, rot, opac, ...}]}]}
  midias/        — imagens, GIFs e vídeos COPIADOS para dentro (o projeto anda sozinho); no JSON viram 'midias/<nome>'
  previa.png     — miniatura do 1º slide
Ao abrir, as mídias saem em <raiz de dados>/Slider/<hash> e os caminhos voltam a ser absolutos.

Exportar: o front fotografa cada slide (o próprio DOM → SVG foreignObject → canvas, igual à tela) e manda quadro a
quadro (`quadro`) para uma pasta de trabalho; `exportar` junta em PDF, PNGs, MP4 (ffmpeg, com transições xfade) ou
PPTX. PPTX sai EDITÁVEL (texto, formas, imagens, vídeos e notas como objetos do PowerPoint, `pptx_nativo`) ou como
imagens (fiel ao pixel). Importar PPTX (`importar_pptx`): textos, formas, imagens, fundo e notas viram camadas.
Guia: Instructions/slider-kanivete.md
"""
import base64, hashlib, html as _html, json, os, re, shutil, subprocess, tempfile, uuid, zipfile

from Functions.midia import raiz_dados

EXTENSAO = ".dknv"
VERSAO = 1
PX_EMU = 6350          # 1920 px = 13,333 pol (slide 16:9 do PowerPoint) → 1 px = 0,5 pt = 6350 EMU


# ───────────────────────── projeto ─────────────────────────
def _midias(doc):
    """(dicionário, chave) de cada caminho de mídia dentro do documento"""
    for s in doc.get("slides", []):
        f = s.get("fundo") or {}
        if isinstance(f.get("img"), str) and f["img"]:
            yield f, "img"
        for c in _todas(s.get("camadas", [])):
            if isinstance(c.get("src"), str) and c["src"]:
                yield c, "src"


def _todas(camadas):
    for c in camadas:
        yield c
        if c.get("tipo") == "grupo":
            yield from _todas(c.get("itens", []))


def salvar(caminho, doc, previa=None):
    if not caminho.lower().endswith(EXTENSAO):
        caminho += EXTENSAO
    doc = json.loads(json.dumps(doc))
    copiar = {}
    for obj, k in _midias(doc):
        p = obj[k]
        if p.startswith("midias/") or not os.path.isfile(p):
            continue
        if p not in copiar:
            copiar[p] = f"{len(copiar):03d}_{os.path.basename(p)}"
        obj[k] = "midias/" + copiar[p]
    # mídias que já vinham de um .dknv aberto (pasta de trabalho) continuam dentro
    tmp = caminho + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("documento.json", json.dumps({**doc, "formato": "dknv", "versao": VERSAO}, ensure_ascii=False))
        for orig, nome in copiar.items():
            ext = os.path.splitext(nome)[1].lower()
            z.write(orig, "midias/" + nome, compress_type=zipfile.ZIP_STORED if ext in (".mp4", ".mov", ".webm", ".jpg", ".jpeg", ".png", ".gif", ".webp") else zipfile.ZIP_DEFLATED)
        if previa and "," in str(previa):
            try:
                z.writestr("previa.png", base64.b64decode(previa.split(",", 1)[1]))
            except Exception:
                pass
    os.replace(tmp, caminho)
    return {"success": True, "path": os.path.abspath(caminho), "midias": len(copiar)}


def abrir(caminho):
    if not os.path.isfile(caminho):
        return {"success": False, "error": "arquivo não encontrado"}
    st = os.stat(caminho)
    pasta = os.path.join(raiz_dados(), "Slider", hashlib.sha1(f"{os.path.abspath(caminho)}|{st.st_mtime}".encode()).hexdigest()[:12])
    with zipfile.ZipFile(caminho) as z:
        doc = json.loads(z.read("documento.json").decode("utf-8"))
        if doc.get("formato") != "dknv":
            return {"success": False, "error": "não é um projeto do Slider Kanivete"}
        for n in z.namelist():
            if n.startswith("midias/") and not n.endswith("/"):
                destino = os.path.join(pasta, n)
                if not os.path.exists(destino):
                    os.makedirs(os.path.dirname(destino), exist_ok=True)
                    with z.open(n) as a, open(destino, "wb") as b:
                        shutil.copyfileobj(a, b)
    for obj, k in _midias(doc):
        if obj[k].startswith("midias/"):
            obj[k] = os.path.join(pasta, obj[k]).replace("\\", "/")
    return {"success": True, "doc": doc, "path": os.path.abspath(caminho)}


def gravar_midia(data_url, nome="colado"):
    """imagem colada (Ctrl+V) ou arrastada sem caminho → arquivo em <raiz de dados>/Slider/colados"""
    cab, b64 = data_url.split(",", 1)
    ext = {"image/png": ".png", "image/jpeg": ".jpg", "image/gif": ".gif", "image/webp": ".webp", "video/mp4": ".mp4"}.get(
        cab.split(";")[0].split(":")[-1], ".png")
    pasta = os.path.join(raiz_dados(), "Slider", "colados"); os.makedirs(pasta, exist_ok=True)
    dados = base64.b64decode(b64)
    base = re.sub(r"[^\w-]", "_", nome)[:30]
    caminho = os.path.join(pasta, f"{base}_{hashlib.sha1(dados).hexdigest()[:10]}{ext}")
    if not os.path.isfile(caminho):
        with open(caminho, "wb") as f:
            f.write(dados)
    return {"success": True, "path": caminho.replace("\\", "/")}


# ───────────────────────── exportar ─────────────────────────
def _pasta_job(job):
    p = os.path.join(tempfile.gettempdir(), "kanivete_slider", re.sub(r"[^\w-]", "", str(job))[:40] or "job")
    os.makedirs(p, exist_ok=True)
    return p


def quadro(job, i, data_url):
    """grava o quadro i (PNG/JPEG em data URL) do job de exportação"""
    cab, b64 = data_url.split(",", 1)
    ext = ".jpg" if "jpeg" in cab else ".png"
    caminho = os.path.join(_pasta_job(job), f"slide_{int(i):03d}{ext}")
    with open(caminho, "wb") as f:
        f.write(base64.b64decode(b64))
    return {"success": True, "path": caminho}


def _quadros(job):
    p = _pasta_job(job)
    return sorted(os.path.join(p, f) for f in os.listdir(p) if f.startswith("slide_"))


def _ffmpeg():
    from Functions.midia import ffmpeg
    return ffmpeg()


def exportar(job, formato, caminho, doc=None, opcoes=None):
    """junta os quadros do job: pdf | png (pasta) | mp4 | pptx (imagens) | pptx_nativo (editável, usa doc)"""
    op = opcoes or {}
    qs = _quadros(job) if formato != "pptx_nativo" else []
    try:
        if formato == "pdf":
            from PIL import Image
            ims = [Image.open(q).convert("RGB") for q in qs]
            if not ims:
                return {"success": False, "error": "nenhum slide"}
            ims[0].save(caminho, "PDF", save_all=True, append_images=ims[1:], resolution=float(op.get("dpi", 144)))
        elif formato == "png":
            os.makedirs(caminho, exist_ok=True)
            base = op.get("base") or "slide"
            for k, q in enumerate(qs, 1):
                shutil.copyfile(q, os.path.join(caminho, f"{base}_{k:02d}{os.path.splitext(q)[1]}"))
        elif formato == "mp4":
            _mp4(qs, caminho, doc or {}, op)
        elif formato == "pptx":
            _pptx_imagens(qs, caminho, doc or {})
        elif formato == "pptx_nativo":
            avisos = pptx_nativo(doc, caminho)
            return {"success": True, "path": caminho, "avisos": avisos}
        else:
            return {"success": False, "error": f"formato desconhecido: {formato}"}
    finally:
        if formato != "pptx_nativo":
            shutil.rmtree(_pasta_job(job), ignore_errors=True)
    return {"success": True, "path": caminho}


def _mp4(qs, caminho, doc, op):
    slides = [s for s in doc.get("slides", []) if not s.get("oculto")] or [{}] * len(qs)
    fps = int(op.get("fps", 30))
    durs = [max(0.5, float((slides[k] if k < len(slides) else {}).get("dur") or op.get("dur", 5))) for k in range(len(qs))]
    XF = {"esmaecer": "fade", "deslizar": "slideleft", "empurrar": "slideleft", "zoom": "zoomin", "cobrir": "wipeleft"}
    cmd = [_ffmpeg(), "-y", "-hide_banner", "-loglevel", "error"]
    for q, d in zip(qs, durs):
        cmd += ["-loop", "1", "-framerate", str(fps), "-t", f"{d:.3f}", "-i", q]
    W, H = int(doc.get("largura", 1920)), int(doc.get("altura", 1080))
    pre = "".join(f"[{k}:v]scale={W}:{H}:force_original_aspect_ratio=decrease,pad={W}:{H}:(ow-iw)/2:(oh-ih)/2,setsar=1,fps={fps},format=yuv420p[v{k}];" for k in range(len(qs)))
    if len(qs) == 1:
        filtro, ult = pre, "[v0]"
    else:
        filtro, ult, t = pre, "[v0]", durs[0]
        for k in range(1, len(qs)):
            tr = (slides[k] if k < len(slides) else {}).get("transicao") or {}
            tipo = XF.get(tr.get("tipo"), None)
            td = min(float(tr.get("dur", 0.5)), durs[k - 1] / 2, durs[k] / 2) if tipo else 0.04
            filtro += f"{ult}[v{k}]xfade=transition={tipo or 'fade'}:duration={td:.3f}:offset={t - td:.3f}[x{k}];"
            ult, t = f"[x{k}]", t - td + durs[k]
    cmd += ["-filter_complex", filtro.rstrip(";"), "-map", ult, "-c:v", "libx264", "-preset", "medium", "-crf", "18",
            "-pix_fmt", "yuv420p", "-movflags", "+faststart", caminho]
    r = subprocess.run(cmd, capture_output=True, text=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if r.returncode != 0:
        raise RuntimeError("ffmpeg: " + (r.stderr or "")[-400:])


def _prs(doc):
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(doc.get("largura", 1920)) * PX_EMU), Emu(int(doc.get("altura", 1080)) * PX_EMU)
    return prs


def _pptx_imagens(qs, caminho, doc):
    prs = _prs(doc)
    slides = [s for s in doc.get("slides", []) if not s.get("oculto")]
    for k, q in enumerate(qs):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.shapes.add_picture(q, 0, 0, prs.slide_width, prs.slide_height)
        notas = (slides[k] if k < len(slides) else {}).get("notas")
        if notas:
            sl.notes_slide.notes_text_frame.text = notas
    prs.save(caminho)


# ── PPTX editável ──
def _rgb(cor):
    from pptx.dml.color import RGBColor
    m = re.match(r"#?([0-9a-fA-F]{6})", str(cor or ""))
    if m:
        return RGBColor.from_string(m.group(1).upper()), 1.0
    m = re.match(r"rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?", str(cor or ""))
    if m:
        r, g, b = (int(float(v)) for v in m.groups()[:3])
        return RGBColor(r, g, b), float(m.group(4)) if m.group(4) else 1.0
    return None, 1.0


def _alpha(fill_el, a):
    """transparência do preenchimento (python-pptx não tem API): <a:alpha val=…/> dentro da cor"""
    if a >= 0.999:
        return
    from pptx.oxml.ns import qn
    from lxml import etree
    for clr in fill_el.iter(qn("a:srgbClr")):
        al = etree.SubElement(clr, qn("a:alpha")); al.set("val", str(int(max(0, min(1, a)) * 100000)))


def _runs_html(h):
    """html simples do texto (b, i, u, span color/font-size/font-weight, br, div/p) → parágrafos de trechos"""
    paras, cur = [[]], []
    pilha = [{}]
    for tok in re.split(r"(<[^>]+>)", h or ""):
        if not tok:
            continue
        if tok.startswith("<"):
            tag = re.match(r"</?\s*(\w+)", tok)
            nome = tag.group(1).lower() if tag else ""
            fecha = tok.startswith("</")
            if nome == "br" or (nome in ("div", "p") and not fecha and any(paras[-1])):
                paras.append([])
                if nome == "br":
                    continue
            if nome in ("div", "p") and fecha:
                continue
            if fecha:
                if len(pilha) > 1:
                    pilha.pop()
                continue
            st = dict(pilha[-1])
            if nome in ("b", "strong"): st["b"] = True
            if nome in ("i", "em"): st["i"] = True
            if nome == "u": st["u"] = True
            m = re.search(r'style="([^"]*)"', tok)
            if m:
                for dec in m.group(1).split(";"):
                    if ":" in dec:
                        k, v = (x.strip().lower() for x in dec.split(":", 1))
                        if k == "color": st["cor"] = v
                        elif k == "font-weight": st["b"] = v in ("bold", "600", "700", "800", "900")
                        elif k == "font-style": st["i"] = v == "italic"
                        elif k == "font-size" and v.endswith("px"): st["tam"] = float(v[:-2])
                        elif k == "text-decoration" and "underline" in v: st["u"] = True
            if not tok.endswith("/>"):
                pilha.append(st)
        else:
            paras[-1].append((_html.unescape(tok), dict(pilha[-1])))
    return paras


def pptx_nativo(doc, caminho):
    from pptx.util import Emu, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    prs = _prs(doc)
    E = lambda v: Emu(int(round(float(v) * PX_EMU)))
    FORMAS = {"ret": MSO_SHAPE.RECTANGLE, "elipse": MSO_SHAPE.OVAL, "tri": MSO_SHAPE.ISOSCELES_TRIANGLE, "losango": MSO_SHAPE.DIAMOND,
              "estrela": MSO_SHAPE.STAR_5_POINT, "seta": MSO_SHAPE.RIGHT_ARROW, "balao": MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
              "hexagono": MSO_SHAPE.HEXAGON, "coracao": MSO_SHAPE.HEART}
    ALIN = {"esq": PP_ALIGN.LEFT, "centro": PP_ALIGN.CENTER, "dir": PP_ALIGN.RIGHT, "just": PP_ALIGN.JUSTIFY}
    VAL = {"topo": MSO_ANCHOR.TOP, "meio": MSO_ANCHOR.MIDDLE, "base": MSO_ANCHOR.BOTTOM}
    avisos = []
    for s in doc.get("slides", []):
        if s.get("oculto"):
            continue
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        f = s.get("fundo") or {}
        if f.get("tipo") == "imagem" and f.get("img") and os.path.isfile(f["img"]):
            sl.shapes.add_picture(f["img"], 0, 0, prs.slide_width, prs.slide_height)
        elif f.get("tipo") == "degrade":
            fill = sl.background.fill; fill.gradient(); fill.gradient_angle = float(f.get("angulo", 90))
            st = fill.gradient_stops
            for k, cor in enumerate((f.get("cor", "#ffffff"), f.get("cor2", "#000000"))):
                c, _ = _rgb(cor)
                if c is not None:
                    st[k].color.rgb = c
        else:
            c, _ = _rgb(f.get("cor", "#ffffff"))
            if c is not None:
                sl.background.fill.solid(); sl.background.fill.fore_color.rgb = c
        for c in _todas(s.get("camadas", [])):
            if c.get("oculto") or c.get("tipo") == "grupo":
                continue
            x, y, w, h = (E(c.get(k, 0)) for k in ("x", "y", "w", "h"))
            sh = None
            t = c.get("tipo")
            if t == "texto":
                sh = sl.shapes.add_textbox(x, y, w, h)
                tf = sh.text_frame; tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = VAL.get(c.get("valin", "topo"), MSO_ANCHOR.TOP)
                cor0, _ = _rgb(c.get("cor", "#000000"))
                for k, par in enumerate(_runs_html(c.get("html") or _html.escape(c.get("texto", "")))):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    p.alignment = ALIN.get(c.get("alin", "esq"))
                    if c.get("entrelinha"):
                        p.line_spacing = float(c["entrelinha"])
                    for txt, st in par:
                        r = p.add_run(); r.text = txt.upper() if c.get("caixa") == "alta" else txt
                        r.font.name = c.get("fonte", "Inter"); r.font.size = Pt(float(st.get("tam", c.get("tam", 40))) * 0.5)
                        r.font.bold = st.get("b", int(c.get("peso", 400)) >= 600); r.font.italic = st.get("i", bool(c.get("italico")))
                        r.font.underline = st.get("u", bool(c.get("sublinhado")))
                        cr, _ = _rgb(st.get("cor")) if st.get("cor") else (cor0, 1)
                        if cr is not None:
                            r.font.color.rgb = cr
            elif t == "forma":
                tipo = c.get("forma", "ret")
                if tipo == "linha":
                    sh = sl.shapes.add_connector(1, x, y + E(c.get("h", 0)) // 2, x + w, y + E(c.get("h", 0)) // 2)
                    cr, _ = _rgb(c.get("traco") or c.get("preench") or "#000000")
                    if cr is not None:
                        sh.line.color.rgb = cr
                    sh.line.width = Pt(max(0.5, float(c.get("espessura", 4)) * 0.5))
                else:
                    mso = FORMAS.get(tipo, MSO_SHAPE.RECTANGLE)
                    if tipo == "ret" and float(c.get("raio", 0)) > 0:
                        mso = MSO_SHAPE.ROUNDED_RECTANGLE
                    sh = sl.shapes.add_shape(mso, x, y, w, h)
                    if mso == MSO_SHAPE.ROUNDED_RECTANGLE:
                        sh.adjustments[0] = min(0.5, float(c["raio"]) / max(1, min(float(c.get("w", 1)), float(c.get("h", 1)))))
                    cr, a = _rgb(c.get("preench"))
                    if cr is not None and c.get("preench") not in ("none", "transparent"):
                        sh.fill.solid(); sh.fill.fore_color.rgb = cr; _alpha(sh.fill._xPr, a * float(c.get("opac", 1)))
                    else:
                        sh.fill.background()
                    tr, _ = _rgb(c.get("traco")) if c.get("traco") and float(c.get("espessura", 0)) > 0 else (None, 1)
                    if tr is not None:
                        sh.line.color.rgb = tr; sh.line.width = Pt(float(c.get("espessura", 2)) * 0.5)
                    else:
                        sh.line.fill.background()
                    sh.shadow.inherit = False
            elif t == "imagem" and c.get("src") and os.path.isfile(c["src"]):
                sh = _figura(sl, c, x, y, w, h)
            elif t == "video" and c.get("src") and os.path.isfile(c["src"]):
                poster = _poster(c["src"])
                try:
                    sh = sl.shapes.add_movie(c["src"], x, y, w, h, poster_frame_image=poster, mime_type="video/mp4")
                except Exception as e:
                    avisos.append(f"vídeo como imagem: {os.path.basename(c['src'])} ({e})")
                    if poster:
                        sh = sl.shapes.add_picture(poster, x, y, w, h)
            if sh is not None and float(c.get("rot", 0)):
                sh.rotation = float(c["rot"]) % 360
            if c.get("sombra") or c.get("mescla"):
                avisos.append(f"{c.get('nome') or t}: sombra/mesclagem não vão para o PPTX editável")
        if s.get("notas"):
            sl.notes_slide.notes_text_frame.text = s["notas"]
    prs.save(caminho)
    return sorted(set(avisos))


def _figura(sl, c, x, y, w, h):
    """imagem com o ajuste da camada: cobrir = recorte (crop) como o object-fit: cover"""
    from PIL import Image
    pic = sl.shapes.add_picture(c["src"], x, y, w, h)
    if c.get("ajuste", "cobrir") == "cobrir":
        try:
            iw, ih = Image.open(c["src"]).size
            caixa, img = float(c["w"]) / float(c["h"]), iw / ih
            if img > caixa:
                corte = (1 - caixa / img) / 2; pic.crop_left = pic.crop_right = corte
            elif img < caixa:
                corte = (1 - img / caixa) / 2; pic.crop_top = pic.crop_bottom = corte
        except Exception:
            pass
    return pic


def _poster(video):
    """1º quadro do vídeo em JPEG (capa no PPTX)"""
    saida = os.path.join(tempfile.gettempdir(), "kanivete_slider", "poster_" + hashlib.sha1(video.encode()).hexdigest()[:10] + ".jpg")
    os.makedirs(os.path.dirname(saida), exist_ok=True)
    if not os.path.isfile(saida):
        subprocess.run([_ffmpeg(), "-y", "-loglevel", "error", "-ss", "0.5", "-i", video, "-frames:v", "1", "-q:v", "3", saida],
                       capture_output=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    return saida if os.path.isfile(saida) else None


# ───────────────────────── importar PPTX ─────────────────────────
def importar_pptx(caminho):
    """PPTX → documento do Slider: textos (com negrito/itálico/cor/tamanho), formas, imagens, fundo sólido e notas.
    Coordenadas na escala do slide de 1920 px de largura."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.enum.text import PP_ALIGN
    prs = Presentation(caminho)
    k = 1920 / prs.slide_width
    H = int(round(prs.slide_height * k))
    pasta = os.path.join(raiz_dados(), "Slider", "importado_" + hashlib.sha1(os.path.abspath(caminho).encode()).hexdigest()[:10])
    os.makedirs(pasta, exist_ok=True)
    ALIN = {PP_ALIGN.CENTER: "centro", PP_ALIGN.RIGHT: "dir", PP_ALIGN.JUSTIFY: "just"}
    AUTO = {1: "ret", 5: "ret", 9: "elipse", 7: "tri", 4: "losango", 92: "estrela", 33: "seta", 10: "hexagono", 21: "coracao"}
    relat, slides = [], []

    def cor_de(fmt):
        try:
            if fmt.type is not None and fmt.rgb is not None:
                return "#" + str(fmt.rgb).lower()
        except Exception:
            return None
        return None

    def camadas_de(shapes, dx=0, dy=0):
        out = []
        for sh in shapes:
            try:
                if sh.left is None or sh.width is None:
                    continue
                x, y, w, h = (sh.left + dx) * k, (sh.top + dy) * k, sh.width * k, sh.height * k
            except Exception:
                continue
            base = {"id": uuid.uuid4().hex[:8], "nome": sh.name, "x": round(x, 1), "y": round(y, 1), "w": round(max(w, 4), 1),
                    "h": round(max(h, 4), 1), "rot": round(float(getattr(sh, "rotation", 0) or 0), 2), "opac": 1}
            st = sh.shape_type
            if st == MSO_SHAPE_TYPE.GROUP:
                out += camadas_de(sh.shapes)
                continue
            if st == MSO_SHAPE_TYPE.PICTURE or (hasattr(sh, "image") and st != MSO_SHAPE_TYPE.MEDIA):
                try:
                    img = sh.image
                    nome = os.path.join(pasta, f"{len(os.listdir(pasta)):03d}_{sh.shape_id}.{img.ext}")
                    with open(nome, "wb") as f:
                        f.write(img.blob)
                    out.append({**base, "tipo": "imagem", "src": nome.replace("\\", "/"), "ajuste": "esticar"})
                except Exception as e:
                    relat.append(f"imagem não lida: {sh.name} ({e})")
                continue
            if st == MSO_SHAPE_TYPE.MEDIA:
                relat.append(f"vídeo/áudio não importado: {sh.name}")
                continue
            if getattr(sh, "has_table", False) and sh.has_table:   # tabela → células (forma + texto) num grupo
                g = uuid.uuid4().hex[:8]; tb = sh.table
                cws = [c.width * k for c in tb.columns]; rhs = [r.height * k for r in tb.rows]
                yy = y
                for i, row in enumerate(tb.rows):
                    xx = x
                    for j, cell in enumerate(row.cells):
                        cor = None
                        try:
                            if cell.fill.type == 1:
                                cor = cor_de(cell.fill.fore_color)
                        except Exception:
                            pass
                        cab = i == 0
                        out.append({"id": uuid.uuid4().hex[:8], "nome": f"Célula {i + 1}.{j + 1}", "tipo": "forma", "forma": "ret", "x": round(xx, 1), "y": round(yy, 1),
                                    "w": round(cws[j], 1), "h": round(rhs[i], 1), "rot": 0, "opac": 1, "preench": cor or ("#2f6fed" if cab else "#eef2fb" if i % 2 else "#ffffff"),
                                    "traco": "#d0d7e5", "espessura": 2, "raio": 0, "grupo": g})
                        txt = "<br>".join(_html.escape(p.text) for p in cell.text_frame.paragraphs)
                        fz = next((r.font.size.pt * 2 for p in cell.text_frame.paragraphs for r in p.runs if r.font.size), 32)
                        out.append({"id": uuid.uuid4().hex[:8], "nome": f"Texto {i + 1}.{j + 1}", "tipo": "texto", "html": txt, "x": round(xx + 16, 1), "y": round(yy, 1),
                                    "w": round(cws[j] - 32, 1), "h": round(rhs[i], 1), "rot": 0, "opac": 1, "fonte": "Calibri", "tam": round(fz), "peso": 700 if cab else 400,
                                    "cor": "#ffffff" if cab and not cor else "#1a1a1a", "alin": "esq", "valin": "meio", "entrelinha": 1.15, "grupo": g})
                        xx += cws[j]
                    yy += rhs[i]
                continue
            if st == MSO_SHAPE_TYPE.AUTO_SHAPE and not (sh.has_text_frame and sh.text_frame.text.strip()):
                try:
                    tipo = AUTO.get(int(sh.auto_shape_type), "ret")
                except Exception:
                    tipo = "ret"
                cor = None
                try:
                    if sh.fill.type == 1:
                        cor = cor_de(sh.fill.fore_color)
                except Exception:
                    pass
                out.append({**base, "tipo": "forma", "forma": tipo, "preench": cor or "#d9d9d9", "traco": None, "espessura": 0, "raio": 0})
                continue
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                fundo = None
                try:
                    if sh.fill.type == 1:
                        fundo = cor_de(sh.fill.fore_color)
                except Exception:
                    pass
                if fundo:
                    out.append({**base, "id": uuid.uuid4().hex[:8], "nome": sh.name + " (fundo)", "tipo": "forma", "forma": "ret",
                                "preench": fundo, "traco": None, "espessura": 0, "raio": 0})
                # o que o modelo do PowerPoint dá por herança (tamanho, centralizado, marcador) quando o texto não diz
                ph = None
                try:
                    ph = int(sh.placeholder_format.type) if sh.is_placeholder else None
                except Exception:
                    ph = None
                TIT, CTIT, SUB, CORPO = 1, 3, 4, (2, 7)   # PP_PLACEHOLDER: TITLE, CENTER_TITLE, SUBTITLE, BODY/OBJECT
                tam_ph = {TIT: 88, CTIT: 120, SUB: 48}.get(ph, 56 if ph in CORPO else None)
                partes, tam0, fonte0, alin0, cor0 = [], None, None, "centro" if ph in (CTIT, SUB) else "esq", None
                for p in sh.text_frame.paragraphs:
                    alin0 = ALIN.get(p.alignment, alin0) if p.alignment is not None else alin0
                    linha = ["&nbsp;&nbsp;&nbsp;&nbsp;" * int(p.level or 0) + "• "] if ph in CORPO and p.text.strip() else []
                    for r in p.runs:
                        t = _html.escape(r.text)
                        if not t:
                            continue
                        fz = r.font.size.pt * 2 if r.font.size else None
                        tam0 = tam0 or fz; fonte0 = fonte0 or r.font.name
                        cor = cor_de(r.font.color) if r.font.color else None
                        cor0 = cor0 or cor
                        estilo = []
                        if cor and cor != cor0: estilo.append(f"color:{cor}")
                        if fz and tam0 and abs(fz - tam0) > 0.5: estilo.append(f"font-size:{fz:.0f}px")
                        if estilo: t = f'<span style="{";".join(estilo)}">{t}</span>'
                        if r.font.bold: t = f"<b>{t}</b>"
                        if r.font.italic: t = f"<i>{t}</i>"
                        if r.font.underline: t = f"<u>{t}</u>"
                        linha.append(t)
                    partes.append("".join(linha))
                out.append({**base, "tipo": "texto", "html": "<br>".join(partes), "fonte": fonte0 or "Calibri", "tam": round(tam0 or tam_ph or 36),
                            "peso": 400, "cor": cor0 or "#1a1a1a", "alin": alin0, "valin": "meio" if ph in (TIT, CTIT) else "topo", "entrelinha": 1.15})
                continue
            if st not in (MSO_SHAPE_TYPE.PLACEHOLDER,):
                relat.append(f"não importado: {sh.name} ({st})")
        return out

    for n, s in enumerate(prs.slides, 1):
        fundo = {"tipo": "cor", "cor": "#ffffff"}
        try:
            if s.background.fill.type == 1:
                fundo["cor"] = cor_de(s.background.fill.fore_color) or "#ffffff"
        except Exception:
            pass
        notas = ""
        try:
            if s.has_notes_slide:
                notas = s.notes_slide.notes_text_frame.text
        except Exception:
            pass
        slides.append({"id": uuid.uuid4().hex[:8], "fundo": fundo, "transicao": {"tipo": "nenhuma", "dur": 0.5}, "dur": 5,
                       "notas": notas, "camadas": camadas_de(s.shapes)})
    nome = os.path.splitext(os.path.basename(caminho))[0]
    return {"success": True, "doc": {"nome": nome, "largura": 1920, "altura": H, "slides": slides}, "relatorio": relat}
