"""
ata_kanivete.py — Ata Kanivete (At, verde; Beta): grava reunião do Google Meet e escreve a ATA.

Caminho (teste de 2026-10-17 em D:/kanivete_testes/ata):
  1. Meet: um Chrome à parte (perfil próprio em <dados>/Ata Kanivete/navegador; login uma vez) aberto com a porta de
     depuração. O app injeta frontend/ata/captura.js antes da página carregar (Page.addScriptToEvaluateOnNewDocument):
     ela grava o microfone que o Meet pediu ("eu") e os áudios que chegam ("outros") em Opus 24 kb/s (~10 MB/h por
     faixa) e anota legendas (nome + texto) e participantes. Tudo volta pela ligação Runtime.addBinding "ataEnviar".
     O Meet manda só ~3 áudios mistos (quem fala mais alto), nunca um por pessoa: as vozes são separadas depois.
  2. Processar: Parakeet TAGARELA (legendas.py) nas duas faixas → separar vozes nos "outros" (sherpa-onnx: pyannote
     segmentation 3.0 + wespeaker resnet34, ~33 MB baixados no 1º uso; 97,7% das palavras certas na reunião de teste)
     → nome de cada grupo de voz pelas legendas do Meet (sem legenda: "Pessoa N", renomeável na tela).
  3. ATA: a Kani (kani_maos._llm_json, JSON travado) — reunião longa vai em blocos e depois junta.
  4. PDF no modelo padrão (PyMuPDF Story, HTML → PDF).
Sessão = pasta <dados>/Ata Kanivete/<AAAA-MM-DD_HHMMSS> com eu.webm, outros.webm, eventos.jsonl e ata.json.
"""

import base64
import datetime
import difflib
import json
import os
import re
import shutil
import socket
import subprocess
import threading
import time
import urllib.request

from Functions.midia import app_dir, raiz_dados

NOME = "Ata Kanivete"
COR = "#1FA971"
NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)
MODELOS = {
    "seg": ("https://github.com/k2-fsa/sherpa-onnx/releases/download/speaker-segmentation-models/"
            "sherpa-onnx-pyannote-segmentation-3-0.tar.bz2", "segmentacao.onnx", 7),
    "voz": ("https://github.com/k2-fsa/sherpa-onnx/releases/download/speaker-recongition-models/"
            "wespeaker_en_voxceleb_resnet34_LM.onnx", "voz.onnx", 26),
}
BLOCO_PALAVRAS = 2200     # a Kani tem 8 mil tokens de contexto: reunião maior vai em blocos


def pasta():
    p = os.path.join(raiz_dados(), "Ata Kanivete")
    os.makedirs(p, exist_ok=True)
    return p


def _modelos_dir():
    p = os.path.join(app_dir(), "modelos_ia", "ata")
    os.makedirs(p, exist_ok=True)
    return p


def _ffmpeg():
    from Functions.audio_cutter import ffmpeg_path
    return ffmpeg_path()


# ───────────────────────── modelos de separar vozes ─────────────────────────
def modelos_prontos():
    return all(os.path.isfile(os.path.join(_modelos_dir(), m[1])) for m in MODELOS.values())


def baixar_modelos(prog=lambda p, m: None):
    import tarfile
    import requests
    d = _modelos_dir()
    for i, (chave, (url, nome, mb)) in enumerate(MODELOS.items()):
        dst = os.path.join(d, nome)
        if os.path.isfile(dst):
            continue
        prog(5 + 40 * i, f"Baixando o separador de vozes ({mb} MB)...")
        tmp = dst + ".parte"
        with requests.get(url, stream=True, timeout=60) as r:
            r.raise_for_status()
            with open(tmp, "wb") as f:
                for b in r.iter_content(1 << 20):
                    f.write(b)
        if url.endswith(".tar.bz2"):
            with tarfile.open(tmp, "r:bz2") as t:
                m = next(x for x in t.getmembers() if x.name.endswith("/model.onnx"))
                with t.extractfile(m) as src, open(dst, "wb") as out:
                    shutil.copyfileobj(src, out)
            os.remove(tmp)
        else:
            os.replace(tmp, dst)


# ───────────────────────── Meet: navegador + captura ─────────────────────────
_meet = {"proc": None, "porta": None, "ws": None, "id": 0, "resp": {}, "sessao": None, "ev": None, "gravando": False,
         "t0": 0.0, "kb": {}, "part": [], "eu": "", "on_evento": None, "lock": threading.Lock()}


def navegador():
    cands = [os.path.join(os.environ.get(v, ""), *p) for v in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA")
             for p in (("Google", "Chrome", "Application", "chrome.exe"), ("Microsoft", "Edge", "Application", "msedge.exe"))]
    return next((c for c in cands if os.path.isfile(c)), None)


def _porta_livre():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    p = s.getsockname()[1]
    s.close()
    return p


def _avisar(o):
    cb = _meet["on_evento"]
    if cb:
        try:
            cb(o)
        except Exception:
            pass


def _cdp(metodo, params=None, esperar=True):
    ws = _meet["ws"]
    if not ws:
        raise RuntimeError("o Meet não está aberto")
    with _meet["lock"]:
        _meet["id"] += 1
        i = _meet["id"]
    ev = threading.Event()
    _meet["resp"][i] = [ev, None]
    ws.send(json.dumps({"id": i, "method": metodo, "params": params or {}}))
    if not esperar:
        return None
    ev.wait(15)
    return (_meet["resp"].pop(i, [None, None])[1] or {}).get("result")


def _receber(payload):
    o = json.loads(payload)
    if o.get("ev") == "inicio":          # cada "● Gravar ATA" = uma sessão nova
        sid = datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
        d = os.path.join(pasta(), sid)
        os.makedirs(d, exist_ok=True)
        _meet.update(sessao=d, ev=open(os.path.join(d, "eventos.jsonl"), "a", encoding="utf-8"), gravando=True,
                     t0=time.time(), kb={}, eu=o.get("eu") or "")
        _salvar_doc(d, {"id": sid, "titulo": "Reunião de " + datetime.datetime.now().strftime("%d/%m/%Y %H:%M"),
                        "quando": datetime.datetime.now().isoformat(timespec="seconds"), "origem": "meet",
                        "link": o.get("url", ""), "eu": _meet["eu"], "estado": "gravando"})
        _avisar({"tipo": "gravando", "id": sid})
    d = _meet["sessao"]
    if not d:
        return
    if "faixa" in o:
        b = base64.b64decode(o["b64"])
        with open(os.path.join(d, o["faixa"] + ".webm"), "ab") as f:
            f.write(b)
        _meet["kb"][o["faixa"]] = _meet["kb"].get(o["faixa"], 0) + len(b) / 1024
        return
    _meet["ev"].write(json.dumps(o, ensure_ascii=False) + "\n")
    _meet["ev"].flush()
    if o.get("ev") == "part":
        _meet["part"] = o.get("p") or []
    elif o.get("ev") == "fim":
        _meet["ev"].close()
        doc = abrir(os.path.basename(d))
        doc.update(estado="gravado", duracao=round(o.get("t") or 0, 1), participantes=_nomes_meet(d), eu=_meet["eu"] or doc.get("eu", ""))
        _salvar_doc(d, doc)
        _meet.update(sessao=None, ev=None, gravando=False)
        _avisar({"tipo": "gravado", "id": doc["id"]})


def _ler_cdp(ws):
    try:
        for msg in ws:
            o = json.loads(msg)
            if "id" in o and o["id"] in _meet["resp"]:
                _meet["resp"][o["id"]][1] = o
                _meet["resp"][o["id"]][0].set()
            elif o.get("method") == "Runtime.bindingCalled" and o["params"].get("name") == "ataEnviar":
                try:
                    _receber(o["params"]["payload"])
                except Exception as e:
                    _avisar({"tipo": "erro", "msg": str(e)})
    except Exception:
        pass
    _meet["ws"] = None
    if _meet["gravando"]:           # o navegador fechou no meio: guarda o que veio
        _receber(json.dumps({"ev": "fim", "t": time.time() - _meet["t0"]}))
    _avisar({"tipo": "fechado"})


def meet_abrir(link="", on_evento=None, auto=False):
    """Abre (ou traz) o Chrome do Meet com a captura. link vazio = página inicial do Meet.
    auto = começa a gravar sozinho ao entrar na chamada (para sozinho ao sair, sempre)."""
    from websockets.sync.client import connect
    _meet["on_evento"] = on_evento or _meet["on_evento"]
    url = link.strip() or "https://meet.google.com/"
    if not re.match(r"^https://meet\.google\.com/", url):
        return {"success": False, "error": "cole um link do Google Meet (https://meet.google.com/...)"}
    if _meet["ws"] and _meet["proc"] and _meet["proc"].poll() is None:
        meet_cmd("auto1" if auto else "auto0")
        _cdp("Page.navigate", {"url": url})
        return {"success": True, "ja_aberto": True}
    exe = navegador()
    if not exe:
        return {"success": False, "error": "não achei o Google Chrome nem o Edge neste PC"}
    porta = _porta_livre()
    perfil = os.path.join(pasta(), "navegador")
    _meet["proc"] = subprocess.Popen([exe, f"--remote-debugging-port={porta}", f"--user-data-dir={perfil}", "--no-first-run",
                                      "--no-default-browser-check", "--new-window", "--window-size=1280,860", "about:blank"])
    _meet["porta"] = porta
    alvo = None
    for _ in range(60):
        try:
            with urllib.request.urlopen(f"http://127.0.0.1:{porta}/json", timeout=2) as r:
                alvo = next((t for t in json.loads(r.read()) if t.get("type") == "page"), None)
            if alvo:
                break
        except Exception:
            pass
        time.sleep(0.5)
    if not alvo:
        return {"success": False, "error": "o navegador não respondeu"}
    ws = connect(alvo["webSocketDebuggerUrl"], max_size=None, open_timeout=10)
    _meet["ws"] = ws
    threading.Thread(target=_ler_cdp, args=(ws,), daemon=True).start()
    _cdp("Runtime.enable")
    _cdp("Page.enable")
    _cdp("Runtime.addBinding", {"name": "ataEnviar"})
    _cdp("Page.addScriptToEvaluateOnNewDocument", {"source": f"window.__ataAuto = {json.dumps(bool(auto))};\n" + _captura_js()})
    _cdp("Page.navigate", {"url": url})
    return {"success": True}


def _captura_js():
    import sys
    base = getattr(sys, "_MEIPASS", None) or os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    with open(os.path.join(base, "frontend", "ata", "captura.js"), encoding="utf-8") as f:
        return f.read()


def meet_cmd(acao):
    """gravar | parar | frente | auto1 | auto0 — o mesmo que o botão por cima do Meet; auto = gravar ao entrar na chamada."""
    if not _meet["ws"]:
        return {"success": False, "error": "o Meet não está aberto"}
    expr = {"gravar": "window.__ata && window.__ata.botao(true)", "parar": "window.__ata && window.__ata.botao(false)",
            "frente": "window.focus()", "auto1": "window.__ataAuto = true; window.__ata && (window.__ata.auto = true)",
            "auto0": "window.__ataAuto = false; window.__ata && (window.__ata.auto = false)"}[acao]
    _cdp("Runtime.evaluate", {"expression": expr, "userGesture": True})
    if acao == "frente":
        _cdp("Page.bringToFront")
    return {"success": True}


def meet_estado():
    aberto = bool(_meet["ws"] and _meet["proc"] and _meet["proc"].poll() is None)
    return {"aberto": aberto, "gravando": _meet["gravando"], "segundos": round(time.time() - _meet["t0"]) if _meet["gravando"] else 0,
            "kb": round(sum(_meet["kb"].values())), "participantes": [n for _, n in _nomes_de(_meet["part"])],
            "id": os.path.basename(_meet["sessao"]) if _meet["sessao"] else None, "navegador": bool(navegador())}


def meet_fechar():
    if _meet["gravando"]:
        meet_cmd("parar")
        time.sleep(2.5)
    p = _meet["proc"]
    if p and p.poll() is None:
        p.terminate()
    return {"success": True}


_LIXO = re.compile(r"^[a-z_]+$|^(Reenquadrar|Fixar|Mais opções|Você|Apresentação)", re.I)


def _nomes_de(part):
    """[(id, nome)] dos blocos do Meet (o texto do bloco traz o nome e às vezes o rótulo de ícone: tira os ícones)."""
    out = []
    for pid, txt in part or []:
        nome = next((s.strip() for s in (txt or "").split("|") if s.strip() and not _LIXO.match(s.strip())), "")
        if nome and nome not in [n for _, n in out]:
            out.append((pid, nome))
    return out


def _nomes_meet(d):
    vistos = []
    for o in _eventos(d):
        if o.get("ev") == "part":
            for _, n in _nomes_de(o.get("p")):
                if n not in vistos:
                    vistos.append(n)
    return vistos


def _eventos(d):
    p = os.path.join(d, "eventos.jsonl")
    if not os.path.isfile(p):
        return []
    out = []
    with open(p, encoding="utf-8") as f:
        for linha in f:
            try:
                out.append(json.loads(linha))
            except ValueError:
                pass
    return out


# ───────────────────────── sessões ─────────────────────────
def _salvar_doc(d, doc):
    tmp = os.path.join(d, "ata.json.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
    os.replace(tmp, os.path.join(d, "ata.json"))


def abrir(sid):
    with open(os.path.join(pasta(), sid, "ata.json"), encoding="utf-8") as f:
        return json.load(f)


def salvar(sid, doc):
    d = os.path.join(pasta(), sid)
    if not os.path.isdir(d):
        return {"success": False, "error": "reunião não encontrada"}
    doc["id"] = sid
    _salvar_doc(d, doc)
    return {"success": True}


_lista_cache = {}        # sid → (data do ata.json, resumo): a lista não relê transcrições longas a cada volta à aba
_em_processo = set()     # sessões sendo processadas agora (o resto marcado "processando" ficou de um app que caiu)


def _resumo(doc):
    tar = (doc.get("ata") or {}).get("tarefas") or []
    r = {k: doc.get(k) for k in ("id", "titulo", "quando", "duracao", "estado", "participantes")}
    r["abertas"] = sum(1 for t in tar if not t.get("feito"))
    return r


def _recuperar(sid, d, doc):
    """Gravação ou processamento que não terminou (o app, o navegador ou o PC caiu no meio): volta a "gravado" com o que
    chegou ao disco, para gerar a ATA. True = mudou."""
    atual = _meet["sessao"] and os.path.basename(_meet["sessao"]) == sid
    if doc.get("estado") == "gravando" and not atual:
        ts = [o.get("t") or 0 for o in _eventos(d)]
        doc.update(estado="gravado", duracao=round(max(ts or [0]), 1) or _duracao(os.path.join(d, "outros.webm")),
                   participantes=_nomes_meet(d), interrompida=True)
    elif doc.get("estado") == "processando" and sid not in _em_processo:
        doc["estado"] = "pronta" if doc.get("ata") else "gravado"
    else:
        return False
    _salvar_doc(d, doc)
    return True


def listar():
    out = []
    for sid in sorted(os.listdir(pasta()), reverse=True):
        d = os.path.join(pasta(), sid)
        p = os.path.join(d, "ata.json")
        try:
            mt = os.path.getmtime(p)
            c = _lista_cache.get(sid)
            if not c or c[0] != mt:
                doc = abrir(sid)
                if _recuperar(sid, d, doc):
                    mt = os.path.getmtime(p)
                c = _lista_cache[sid] = (mt, _resumo(doc))
        except Exception:
            continue
        out.append(c[1])
    return out


def _sem_acento(s):
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", str(s or "").lower()) if unicodedata.category(c) != "Mn")


def buscar(termo):
    """Reuniões com o termo no título, na ATA ou na transcrição (sem diferença de acento/maiúscula):
    [resumo + trecho (a frase onde aparece) + t (segundo da fala, se veio da transcrição)]."""
    q = _sem_acento(termo).strip()
    if len(q) < 2:
        return []
    out = []
    for r in listar():
        try:
            doc = abrir(r["id"])
        except Exception:
            continue
        achou, trecho, t = q in _sem_acento(doc.get("titulo")), "", None
        for tu in doc.get("turnos") or []:
            tx = _sem_acento(tu.get("texto"))
            i = tx.find(q)
            if i >= 0:
                orig = tu["texto"]
                a = max(0, i - 40)
                trecho = ("…" if a else "") + orig[a:i + len(q) + 60].strip() + ("…" if i + len(q) + 60 < len(orig) else "")
                t, achou = tu.get("ini"), True
                break
        if not achou and q in _sem_acento(json.dumps(doc.get("ata") or {}, ensure_ascii=False)):
            achou, trecho = True, "na ATA"
        if achou:
            out.append(dict(r, trecho=trecho, t=t))
    return out


def tarefas_abertas():
    """Tarefas ainda não feitas de todas as reuniões (mais recentes primeiro)."""
    out = []
    for r in listar():
        if not r.get("abertas"):
            continue
        try:
            doc = abrir(r["id"])
        except Exception:
            continue
        for i, t in enumerate((doc.get("ata") or {}).get("tarefas") or []):
            if not t.get("feito"):
                out.append(dict(t, id=r["id"], i=i, reuniao=doc.get("titulo"), quando=doc.get("quando")))
    return out


def marcar_tarefa(sid, i, feito=True):
    doc = abrir(sid)
    tar = (doc.get("ata") or {}).get("tarefas") or []
    if not 0 <= int(i) < len(tar):
        return {"success": False, "error": "tarefa não encontrada"}
    tar[int(i)]["feito"] = bool(feito)
    return salvar(sid, doc)


def audio_mix(sid):
    """Uma faixa só (a sua voz + os outros) com índice para pular no tempo: o .webm do MediaRecorder não tem índice e o
    player não vai direto a um ponto de uma reunião longa. Refeita só se uma faixa mudar."""
    d = os.path.join(pasta(), sid)
    dst = os.path.join(d, "mix.webm")
    fx = [os.path.join(d, f) for f in ("eu.webm", "outros.webm") if os.path.isfile(os.path.join(d, f))]
    if not fx:
        return None
    if os.path.isfile(dst) and os.path.getmtime(dst) >= max(map(os.path.getmtime, fx)):
        return dst
    cmd = [_ffmpeg(), "-y", "-v", "error"]
    for f in fx:
        cmd += ["-i", f]
    if len(fx) > 1:
        cmd += ["-filter_complex", "amix=inputs=2:duration=longest:normalize=0"]
    tmp = os.path.join(d, "_mix.webm")
    cmd += ["-ac", "1", "-c:a", "libopus", "-b:a", "32k", "-application", "voip", tmp]
    subprocess.run(cmd, check=True, capture_output=True, creationflags=NO_WINDOW)
    os.replace(tmp, dst)
    return dst


def apagar(sid):
    d = os.path.join(pasta(), sid)
    if os.path.isdir(d) and os.path.isfile(os.path.join(d, "ata.json")):
        shutil.rmtree(d, ignore_errors=True)
    return {"success": True}


def importar_audio(caminho):
    """Arquivo de áudio/vídeo de uma reunião → sessão nova (vozes todas na faixa "outros")."""
    sid = datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
    d = os.path.join(pasta(), sid)
    os.makedirs(d, exist_ok=True)
    r = subprocess.run([_ffmpeg(), "-y", "-v", "error", "-i", caminho, "-vn", "-ac", "1", "-ar", "16000", "-c:a", "libopus",
                        "-b:a", "24k", "-application", "voip", os.path.join(d, "outros.webm")],
                       capture_output=True, text=True, creationflags=NO_WINDOW)
    if r.returncode != 0:
        shutil.rmtree(d, ignore_errors=True)
        return {"success": False, "error": "não consegui ler o áudio: " + (r.stderr or "")[-200:]}
    dur = _duracao(os.path.join(d, "outros.webm"))
    nome = os.path.splitext(os.path.basename(caminho))[0]
    _salvar_doc(d, {"id": sid, "titulo": nome, "quando": datetime.datetime.now().isoformat(timespec="seconds"),
                    "origem": "arquivo", "arquivo": caminho, "eu": "", "estado": "gravado", "duracao": dur, "participantes": []})
    return {"success": True, "id": sid}


def _duracao(p):
    from Functions.audio_cutter import ffmpeg_path
    probe = ffmpeg_path().replace("ffmpeg.exe", "ffprobe.exe")
    try:
        r = subprocess.run([probe, "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", p],
                           capture_output=True, text=True, creationflags=NO_WINDOW)
        return round(float(r.stdout.strip()), 1)
    except Exception:
        return 0


# ───────────────────────── processar: transcrever, separar vozes, nomes ─────────────────────────
def _wav16(src, dst):
    subprocess.run([_ffmpeg(), "-y", "-v", "error", "-i", src, "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le", dst],
                   check=True, creationflags=NO_WINDOW)


_asr = {}
# trechos mais curtos que os das legendas (20 s): o Parakeet às vezes devolve VAZIO num trecho longo e o corte forçado
# aos 20 s parte frase ("cinco mil" → "cinco anos. Mil"); pausa mínima menor = corte cai entre frases
VAD_REUNIAO = {"min_silence_duration_ms": 350, "speech_pad_ms": 150, "max_speech_duration_s": 12}


def _transcrever(wav, stop=None):
    import gc
    from Functions import legendas
    if not legendas.modelo_pronto("pt"):
        legendas.preparar_modelo("pt", None, stop)
    if "m" not in _asr:
        import onnx_asr
        vad = onnx_asr.load_vad("silero", legendas._pasta_modelos())
        m = onnx_asr.load_model(legendas.MODELOS["pt"]["tipo"], legendas._pasta("pt"), quantization="int8")
        _asr["m"] = m.with_vad(vad, **VAD_REUNIAO).with_timestamps()
    palavras = []
    for ini, x in legendas._pedacos(wav):
        for seg in _asr["m"].recognize(x, sample_rate=16000):
            if stop is not None and stop.is_set():
                raise InterruptedError
            seg.start, seg.end = seg.start + ini, seg.end + ini
            palavras += legendas._palavras(seg)
        del x
        gc.collect()
    return legendas.corrigir(palavras)


def liberar():
    """Solta o reconhecimento da memória (chamado ao terminar de processar)."""
    _asr.clear()


def _glossario_aplicar(palavras, glossario):
    """Termo do glossário que o reconhecimento escreveu parecido ("Canivete" → "Kanivete") volta certo."""
    termos = [t.strip() for t in glossario if t.strip() and " " not in t.strip() and len(t.strip()) >= 4]
    for w in palavras:
        nua = w[2].strip(".,;:!?…\"'()")
        for t in termos:
            if nua.lower() != t.lower() and difflib.SequenceMatcher(None, nua.lower(), t.lower()).ratio() >= 0.78:
                w[2] = w[2].replace(nua, t)
                break
    return palavras


def _separar(wav, prog):
    import numpy as np
    import sherpa_onnx
    import soundfile as sf
    if not modelos_prontos():
        baixar_modelos(prog)
    d = _modelos_dir()
    cfg = sherpa_onnx.OfflineSpeakerDiarizationConfig(
        segmentation=sherpa_onnx.OfflineSpeakerSegmentationModelConfig(
            pyannote=sherpa_onnx.OfflineSpeakerSegmentationPyannoteModelConfig(model=os.path.join(d, "segmentacao.onnx")),
            num_threads=2),
        embedding=sherpa_onnx.SpeakerEmbeddingExtractorConfig(model=os.path.join(d, "voz.onnx"), num_threads=2),
        clustering=sherpa_onnx.FastClusteringConfig(num_clusters=-1, threshold=0.5),
        min_duration_on=0.3, min_duration_off=0.5)
    sd = sherpa_onnx.OfflineSpeakerDiarization(cfg)
    x, _ = sf.read(wav, dtype="float32")
    if x.ndim > 1:
        x = x.mean(1)
    r = sd.process(np.ascontiguousarray(x)).sort_by_start_time()
    return [{"ini": s.start, "fim": s.end, "spk": int(s.speaker)} for s in r]


def _legendas_falas(eventos, nomes):
    """Intervalos [ini, fim, nome] de quem a legenda do Meet mostrava falando (o nome é a linha que bate com um participante)."""
    out = []
    for o in eventos:
        if o.get("ev") != "cc":
            continue
        quem = None
        for linha in (o.get("txt") or "").split("\n"):
            linha = linha.strip()
            if linha in nomes or linha in ("Você", "You"):
                quem = linha
        if quem and quem not in ("Você", "You"):
            out.append([o["t"] - 1.5, o["t"] + 0.3, quem])
    return out


def _nome_no(intervalos, t):
    for a, b, n in intervalos:
        if a <= t <= b:
            return n
    return None


def processar(sid, prog=lambda p, m: None, stop=None):
    d = os.path.join(pasta(), sid)
    doc = abrir(sid)
    doc["estado"] = "processando"
    doc.pop("erro", None)
    _em_processo.add(sid)
    _salvar_doc(d, doc)
    try:
        palavras = []           # [ini, fim, texto, spk]
        eu_webm, out_webm = os.path.join(d, "eu.webm"), os.path.join(d, "outros.webm")
        glossario = doc.get("glossario") or []
        if os.path.isfile(eu_webm):
            prog(5, "Transcrevendo a sua voz...")
            w = os.path.join(d, "_eu16.wav")
            _wav16(eu_webm, w)
            palavras += [p + ["eu"] for p in _transcrever(w, stop)]
            os.remove(w)
        falantes = {"eu": doc.get("eu") or "Eu"} if palavras else {}
        if os.path.isfile(out_webm):
            prog(25, "Transcrevendo os outros participantes...")
            w = os.path.join(d, "_outros16.wav")
            _wav16(out_webm, w)
            ow = _transcrever(w, stop)
            if ow:
                prog(50, "Separando as vozes...")
                segs = _separar(w, prog)
                nomes = [n for n in (doc.get("participantes") or []) if n != doc.get("eu")]
                cc = _legendas_falas(_eventos(d), nomes)
                # nome por grupo de voz: o que a legenda mais mostrava enquanto ele falava
                votos = {}
                for s in segs:
                    t = s["ini"]
                    while t < s["fim"]:
                        n = _nome_no(cc, t)
                        if n:
                            votos.setdefault(s["spk"], {}).setdefault(n, 0)
                            votos[s["spk"]][n] += 1
                        t += 0.1
                ordem = sorted({s["spk"] for s in segs}, key=lambda k: next(s["ini"] for s in segs if s["spk"] == k))
                for i, k in enumerate(ordem):
                    falantes[f"p{k}"] = max(votos[k], key=votos[k].get) if k in votos else f"Pessoa {i + 1}"
                for p in ow:
                    m = (p[0] + p[1]) / 2
                    dentro = [s for s in segs if s["ini"] - 0.2 <= m <= s["fim"] + 0.2] or segs
                    s = min(dentro, key=lambda s: 0 if s["ini"] <= m <= s["fim"] else min(abs(m - s["ini"]), abs(m - s["fim"])))
                    palavras.append(p + [f"p{s['spk']}" if segs else "p0"])
                if not segs:
                    falantes["p0"] = "Pessoa 1"
            os.remove(w)
        if stop is not None and stop.is_set():
            raise InterruptedError
        _glossario_aplicar(palavras, glossario)
        palavras.sort(key=lambda p: p[0])
        turnos = []
        for p in palavras:
            if turnos and turnos[-1]["spk"] == p[3] and p[0] - turnos[-1]["fim"] < 2.5:
                turnos[-1]["texto"] += " " + p[2]
                turnos[-1]["fim"] = round(p[1], 2)
            else:
                turnos.append({"spk": p[3], "ini": round(p[0], 2), "fim": round(p[1], 2), "texto": p[2]})
        doc.update(turnos=turnos, falantes=falantes)
        if not turnos:
            doc["estado"] = "gravado"
            _salvar_doc(d, doc)
            return {"success": False, "error": "não encontrei fala na gravação"}
        _salvar_doc(d, doc)      # transcrição guardada antes da Kani: cancelar ou falhar daqui em diante não a perde
        try:
            prog(66, "Preparando o áudio para ouvir...")
            audio_mix(sid)
        except Exception:
            pass
        prog(70, "A Kani está escrevendo a ATA...")
        doc["ata"] = escrever_ata(doc, prog, stop)
        if doc["ata"].get("titulo"):
            doc["titulo"] = doc["ata"]["titulo"]
        doc["estado"] = "pronta"
        _salvar_doc(d, doc)
        prog(100, "ATA pronta")
        return {"success": True, "doc": doc}
    except InterruptedError:
        doc["estado"] = "pronta" if doc.get("ata") else "gravado"
        _salvar_doc(d, doc)
        return {"success": False, "cancelled": True}
    except Exception as e:
        doc["estado"] = "pronta" if doc.get("ata") else "gravado"
        doc["erro"] = str(e)
        _salvar_doc(d, doc)
        raise
    finally:
        _em_processo.discard(sid)


# ───────────────────────── ATA pela Kani ─────────────────────────
SISTEMA = ("Você é secretária de reuniões. Escreva a ATA da reunião a partir da transcrição, em português do Brasil, formal e "
           "objetiva. Use SÓ o que foi dito: não invente nomes, valores, datas ou decisões. Decisão = o que o grupo fechou. "
           "Tarefa = alguém ficou de fazer algo (responsável, o quê, prazo como foi dito; sem prazo: \"a definir\"). "
           "Pendência = ficou em aberto ou depende de algo. Corrija erros óbvios da transcrição pelo contexto, mas nunca "
           "troque um termo do glossário por outro nome.")
# estilo escolhido na tela (doc["estilo"]) + pedido livre do usuário (doc["pedido"]) vão no fim das instruções
ESTILOS = {
    "formal": ("Formal", ""),
    "resumida": ("Resumida", " Seja breve: resumo em até 3 frases, no máximo 4 temas na discussão, itens curtos."),
    "detalhada": ("Detalhada", " Seja detalhada: registre cada argumento importante, números citados e quem defendeu cada "
                               "posição na discussão."),
    "cliente": ("Reunião com cliente", " É uma reunião com cliente: na discussão destaque as necessidades do cliente, a "
                                       "proposta, valores e objeções; nas tarefas, os próximos passos de cada lado."),
    "equipe": ("Alinhamento de equipe", " É um alinhamento de equipe: na discussão, o andamento de cada frente e os "
                                        "bloqueios; nas tarefas, quem destrava o quê."),
    "aula": ("Aula ou treinamento", " É uma aula ou treinamento: a discussão vira os tópicos ensinados com os pontos-chave "
                                    "de cada um; decisões = combinados; tarefas = exercícios ou entregas pedidas."),
}
SCHEMA = {"type": "object", "required": ["titulo", "pauta", "resumo", "discussao", "decisoes", "tarefas", "pendencias", "proxima"],
          "properties": {
              "titulo": {"type": "string"},
              "pauta": {"type": "array", "items": {"type": "string"}, "maxItems": 10},
              "resumo": {"type": "string"},
              "discussao": {"type": "array", "maxItems": 10, "items": {"type": "object", "required": ["tema", "texto"],
                            "properties": {"tema": {"type": "string"}, "texto": {"type": "string"}}}},
              "decisoes": {"type": "array", "items": {"type": "string"}},
              "tarefas": {"type": "array", "items": {"type": "object", "required": ["responsavel", "tarefa", "prazo"],
                          "properties": {"responsavel": {"type": "string"}, "tarefa": {"type": "string"},
                                         "prazo": {"type": "string"}}}},
              "pendencias": {"type": "array", "items": {"type": "string"}},
              "proxima": {"type": "string"}}}


def transcricao_texto(doc):
    """[mm:ss] Nome: fala por turno; turnos seguidos da mesma pessoa (grupos de voz renomeados iguais) viram um só."""
    fal = doc.get("falantes") or {}
    linhas = []
    for t in doc.get("turnos") or []:
        nome = fal.get(t["spk"], t["spk"])
        if linhas and linhas[-1][1] == nome and t["ini"] - linhas[-1][2] < 2.5:
            linhas[-1][3] += " " + t["texto"]
            linhas[-1][2] = t["fim"]
        else:
            linhas.append([t["ini"], nome, t["fim"], t["texto"]])
    return "\n".join(f"[{int(i // 60):02d}:{int(i % 60):02d}] {n}: {tx}" for i, n, _, tx in linhas)


def escrever_ata(doc, prog=lambda p, m: None, stop=None):
    from Functions import kani
    from Functions.kani_maos import _llm_json
    if not kani.estado().get("pronto"):
        raise RuntimeError("a Kani ainda não está instalada: abra a Kani (barra lateral › IA › Kani) uma vez para baixar o modelo")

    def parou():
        if stop is not None and stop.is_set():
            raise InterruptedError
    sistema = SISTEMA + ESTILOS.get(doc.get("estilo") or "formal", ESTILOS["formal"])[1]
    if (doc.get("pedido") or "").strip():
        sistema += " Pedido do usuário para esta ATA (siga, sem inventar fatos): " + doc["pedido"].strip()
    fal = doc.get("falantes") or {}
    cab = f"Participantes: {', '.join(dict.fromkeys(fal.values()))}\n"
    if doc.get("glossario"):
        cab += f"Glossário (escreva exatamente assim): {', '.join(doc['glossario'])}\n"
    linhas = transcricao_texto(doc).split("\n")
    blocos, atual, n = [], [], 0
    for l in linhas:
        atual.append(l)
        n += len(l.split())
        if n >= BLOCO_PALAVRAS:
            blocos.append(atual)
            atual, n = [], 0
    if atual:
        blocos.append(atual)
    if len(blocos) == 1:
        return _llm_json(sistema, cab + "\nTranscrição:\n" + "\n".join(blocos[0]), SCHEMA, 2000)
    # reunião longa: uma ATA parcial por bloco, depois a Kani junta as parciais numa só
    parciais = []
    for i, b in enumerate(blocos):
        parou()
        prog(70 + 20 * i / len(blocos), f"A Kani está lendo a parte {i + 1} de {len(blocos)}...")
        parciais.append(_llm_json(sistema, cab + f"\nTranscrição (parte {i + 1} de {len(blocos)}):\n" + "\n".join(b), SCHEMA, 1500))
    parou()
    prog(92, "Juntando as partes da ATA...")
    return _llm_json(sistema + " Agora você recebe ATAs parciais de partes seguidas da MESMA reunião: junte numa ATA só, "
                     "sem repetir itens.", cab + "\nATAs parciais:\n" + json.dumps(parciais, ensure_ascii=False), SCHEMA, 2000)


# ───────────────────────── PDF no modelo padrão ─────────────────────────
def _esc(s):
    return (str(s or "")).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _data_br(iso):
    try:
        return datetime.datetime.fromisoformat(iso).strftime("%d/%m/%Y")
    except Exception:
        return ""


def _hora(iso):
    try:
        return datetime.datetime.fromisoformat(iso).strftime("%H:%M")
    except Exception:
        return ""


def _dur(s):
    s = int(s or 0)
    return f"{s // 3600} h {s % 3600 // 60:02d} min" if s >= 3600 else f"{s // 60} min {s % 60:02d} s"


def html_ata(doc):
    a = doc.get("ata") or {}
    part = list(dict.fromkeys((doc.get("falantes") or {}).values())) or doc.get("participantes") or []
    lista = lambda xs: "<ul>" + "".join(f"<li>{_esc(x)}</li>" for x in xs if str(x).strip()) + "</ul>" if xs else ""
    h = [f'<p class="rot">ATA DE REUNIÃO</p><h1>{_esc(a.get("titulo") or doc.get("titulo"))}</h1>',
         '<table class="meta"><tr><th>DATA</th><th>HORÁRIO</th><th>DURAÇÃO</th><th>PARTICIPANTES</th></tr>',
         f'<tr><td>{_data_br(doc.get("quando"))}</td><td>{_hora(doc.get("quando"))}</td><td>{_dur(doc.get("duracao"))}</td>'
         f'<td>{_esc(", ".join(part))}</td></tr></table>']
    if a.get("pauta"):
        h.append("<h2>PAUTA</h2>" + lista(a["pauta"]))
    if a.get("resumo"):
        h.append(f"<h2>RESUMO</h2><p>{_esc(a['resumo'])}</p>")
    if a.get("discussao"):
        h.append("<h2>DISCUSSÃO</h2>" + "".join(f"<h3>{_esc(x.get('tema'))}</h3><p>{_esc(x.get('texto'))}</p>" for x in a["discussao"]))
    if a.get("decisoes"):
        h.append("<h2>DECISÕES</h2>" + lista(a["decisoes"]))
    if a.get("tarefas"):
        h.append('<h2>TAREFAS</h2><table class="tar"><tr><th>RESPONSÁVEL</th><th>TAREFA</th><th>PRAZO</th></tr>' +
                 "".join(f"<tr><td>{_esc(t.get('responsavel'))}</td><td>{_esc(t.get('tarefa'))}</td><td>{_esc(t.get('prazo'))}"
                         f"{' (concluída)' if t.get('feito') else ''}</td></tr>" for t in a["tarefas"]) + "</table>")
    if a.get("pendencias"):
        h.append("<h2>PENDÊNCIAS</h2>" + lista(a["pendencias"]))
    if a.get("proxima"):
        h.append(f"<h2>PRÓXIMA REUNIÃO</h2><p>{_esc(a['proxima'])}</p>")
    return "".join(h)


CSS_PDF = f"""
@font-face {{ font-family: tx; src: url(segoeui.ttf); }}
@font-face {{ font-family: tx; src: url(segoeuib.ttf); font-weight: bold; }}
body {{ font-family: tx, sans-serif; font-size: 10pt; color: #1F2328; line-height: 1.5; }}
.rot {{ color: {COR}; font-size: 7.5pt; font-weight: bold; margin: 0 0 2pt 0; letter-spacing: 1pt; }}
h1 {{ font-size: 19pt; font-weight: bold; margin: 0 0 10pt 0; line-height: 1.2; }}
h2 {{ color: {COR}; font-size: 8.5pt; font-weight: bold; margin: 14pt 0 4pt 0; letter-spacing: 0.6pt; }}
h3 {{ font-size: 10pt; font-weight: bold; margin: 6pt 0 1pt 0; }}
p {{ margin: 0 0 4pt 0; }}
ul {{ margin: 0 0 0 0; padding-left: 14pt; }}
li {{ margin: 0 0 2pt 0; }}
table {{ width: 100%; border-collapse: collapse; }}
th {{ color: #6B7280; font-size: 7.5pt; font-weight: bold; text-align: left; padding: 4pt 6pt 4pt 0; }}
td {{ padding: 4pt 6pt 4pt 0; vertical-align: top; }}
.meta {{ border-top: 1.2pt solid {COR}; border-bottom: 0.5pt solid #E5E7EB; }}
.tar th {{ border-bottom: 0.8pt solid {COR}; }}
.tar td {{ border-bottom: 0.4pt solid #E5E7EB; }}
"""


def exportar_pdf(sid, caminho, doc=None):
    import fitz
    doc = doc or abrir(sid)
    fontes = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    arq = fitz.Archive(fontes) if os.path.isfile(os.path.join(fontes, "segoeui.ttf")) else None
    story = fitz.Story(html=html_ata(doc), user_css=CSS_PDF, archive=arq)
    larg, alt = fitz.paper_size("a4")
    m = 56.7     # 20 mm
    corpo = fitz.Rect(m, m, larg - m, alt - m)
    import io
    buf = io.BytesIO()
    w = fitz.DocumentWriter(buf)
    mais = True
    while mais:
        dev = w.begin_page(fitz.Rect(0, 0, larg, alt))
        mais, _ = story.place(corpo)
        story.draw(dev)
        w.end_page()
    w.close()
    # faixa verde no topo e rodapé "Ata Kanivete · página"
    pdf = fitz.open("pdf", buf.getvalue())
    verde = tuple(int(COR[i:i + 2], 16) / 255 for i in (1, 3, 5))
    for i, pg in enumerate(pdf):
        pg.draw_rect(fitz.Rect(0, 0, larg, 4), color=None, fill=verde)
        pg.insert_text((m, alt - 30), NOME, fontsize=7.5, color=(0.42, 0.45, 0.5))
        pg.insert_text((larg - m - 8, alt - 30), str(i + 1), fontsize=7.5, color=(0.42, 0.45, 0.5))
    pdf.set_metadata({"title": (doc.get("ata") or {}).get("titulo") or doc.get("titulo") or "ATA", "author": NOME})
    pdf.save(caminho)
    pdf.close()
    return {"success": True, "caminho": caminho}


# ───────────────────────── Word (.docx) no mesmo modelo ─────────────────────────
def exportar_docx(sid, caminho, doc=None):
    """A ATA em Word editável, com as mesmas seções e o verde do PDF (python-docx)."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt, RGBColor
    doc = doc or abrir(sid)
    a = doc.get("ata") or {}
    verde = RGBColor.from_string(COR[1:])
    cinza = RGBColor(0x6B, 0x72, 0x80)
    w = Document()
    st = w.styles["Normal"]
    st.font.name, st.font.size = "Segoe UI", Pt(10)

    def par(txt, tam=10, cor=None, negrito=False, antes=0, depois=4):
        p = w.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(antes), Pt(depois)
        r = p.add_run(str(txt or ""))
        r.font.size, r.bold = Pt(tam), negrito
        if cor:
            r.font.color.rgb = cor
        return p

    def secao(nome):
        par(nome, 8.5, verde, True, 14, 4)

    def lista(xs):
        for x in xs:
            if str(x).strip():
                w.add_paragraph(str(x), style="List Bullet")

    part = list(dict.fromkeys((doc.get("falantes") or {}).values())) or doc.get("participantes") or []
    par("ATA DE REUNIÃO", 7.5, verde, True, 0, 2)
    par(a.get("titulo") or doc.get("titulo"), 19, None, True, 0, 8)
    par(f"{_data_br(doc.get('quando'))} · {_hora(doc.get('quando'))} · {_dur(doc.get('duracao'))}"
        + (f" · Participantes: {', '.join(part)}" if part else ""), 9, cinza, False, 0, 6)
    if a.get("pauta"):
        secao("PAUTA")
        lista(a["pauta"])
    if a.get("resumo"):
        secao("RESUMO")
        par(a["resumo"])
    if a.get("discussao"):
        secao("DISCUSSÃO")
        for x in a["discussao"]:
            par(x.get("tema"), 10, None, True, 6, 1)
            par(x.get("texto"))
    if a.get("decisoes"):
        secao("DECISÕES")
        lista(a["decisoes"])
    if a.get("tarefas"):
        secao("TAREFAS")
        tb = w.add_table(rows=1, cols=3)
        tb.alignment = WD_TABLE_ALIGNMENT.LEFT
        tb.style = "Light List Accent 6" if "Light List Accent 6" in [s.name for s in w.styles] else tb.style
        for c, nome in zip(tb.rows[0].cells, ("RESPONSÁVEL", "TAREFA", "PRAZO")):
            c.text = nome
        for t in a["tarefas"]:
            cs = tb.add_row().cells
            cs[0].text, cs[1].text = t.get("responsavel") or "", t.get("tarefa") or ""
            cs[2].text = (t.get("prazo") or "") + (" (concluída)" if t.get("feito") else "")
    if a.get("pendencias"):
        secao("PENDÊNCIAS")
        lista(a["pendencias"])
    if a.get("proxima"):
        secao("PRÓXIMA REUNIÃO")
        par(a["proxima"])
    w.core_properties.title = a.get("titulo") or doc.get("titulo") or "ATA"
    w.core_properties.author = NOME
    w.save(caminho)
    return {"success": True, "caminho": caminho}
